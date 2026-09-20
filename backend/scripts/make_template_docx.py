#!/usr/bin/env python
"""Tạo file .docx mẫu đã đúng thể thức để workflow 3 điền vào.

Điền vào mẫu thay vì dựng từ code có cái lợi lớn: thể thức (phông, cỡ, lề, canh
chỉnh) kế thừa sẵn từ file, không phải set lại từng thuộc tính và không sợ lệch
chuẩn. Chỗ cần thay đánh dấu bằng {{placeholder}}.

    uv run python scripts/make_template_docx.py
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt

OUTPUT = Path("data/templates/bao_cao_tai_nguyen.docx")

# Theo Nghị định 30/2020/NĐ-CP. Đây là quy định cho văn bản hành chính, áp cho
# văn bản workflow 3 sinh ra; bước soát ở workflow 2 không ép mẫu này.
FONT = "Times New Roman"
SIZE_NOI_DUNG = Pt(13)
MARGINS_MM = {"top": 22, "bottom": 22, "left": 32, "right": 17}


def _set_default_style(document: docx.Document) -> None:
    style = document.styles["Normal"]
    style.font.name = FONT
    style.font.size = SIZE_NOI_DUNG
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.3


def _para(document, text, *, bold=False, size=None, align=None, italic=False):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = size or SIZE_NOI_DUNG
    run.font.bold = bold
    run.font.italic = italic
    if align is not None:
        paragraph.alignment = align
    return paragraph


def build() -> Path:
    document = docx.Document()
    _set_default_style(document)

    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)     # A4
    section.top_margin = Mm(MARGINS_MM["top"])
    section.bottom_margin = Mm(MARGINS_MM["bottom"])
    section.left_margin = Mm(MARGINS_MM["left"])
    section.right_margin = Mm(MARGINS_MM["right"])

    # --- thể thức đầu văn bản ---
    _para(document, "{{noi_gui}}", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    _para(document, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(document, "Độc lập - Tự do - Hạnh phúc", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(document, "Số: {{so_ky_hieu}}", align=WD_ALIGN_PARAGRAPH.LEFT)
    _para(document, "{{dia_danh}}, {{ngay_bao_cao}}", italic=True,
          align=WD_ALIGN_PARAGRAPH.RIGHT)

    # --- tên loại và trích yếu ---
    _para(document, "BÁO CÁO", bold=True, size=Pt(14), align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(document, "{{trich_yeu}}", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    _para(document, "Kính gửi: {{noi_nhan}}", bold=True)
    _para(document, "{{can_cu}}", italic=True)

    # Mốc để chèn các mục nội dung (đoạn văn và bảng).
    _para(document, "{{sections}}")

    # --- nơi nhận và chữ ký ---
    _para(document, "Nơi nhận:", bold=True)
    _para(document, "- Như trên;")
    _para(document, "- Lưu: VT.")
    _para(document, "{{chuc_vu_ky}}", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _para(document, "{{nguoi_ky}}", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"Đã tạo {path}")
