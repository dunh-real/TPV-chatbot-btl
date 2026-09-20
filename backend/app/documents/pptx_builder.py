"""Kết xuất bộ slide từ JSON trung gian - hoàn toàn bằng code.

LLM chỉ sinh ra cấu trúc (`DeckSpec`): slide nào, tiêu đề gì, gạch đầu dòng gì.
Việc đặt shape, chọn layout, canh lề, chèn ảnh đều do đây làm. Để model thao tác
trực tiếp lên PowerPoint thì không kiểm soát được kết quả và cũng không test được.

Năm kiểu slide, cố định - outline do LLM sinh phải nằm trong danh sách này:
    title    bìa
    summary  các chỉ tiêu chính dạng ô số lớn
    bullet   gạch đầu dòng
    chart    một biểu đồ PNG (do code vẽ) kèm chú thích
    table    bảng số liệu
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field, replace
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

logger = logging.getLogger(__name__)

SLIDE_KINDS = ("title", "summary", "bullet", "chart", "table")

FONT = "Arial"                 # slide chiếu: sans-serif dễ đọc hơn Times
COLOR_TITLE = RGBColor(0x0F, 0x24, 0x47)
COLOR_TEXT = RGBColor(0x1F, 0x2A, 0x37)
COLOR_ACCENT = RGBColor(0x25, 0x63, 0xEB)
COLOR_MUTED = RGBColor(0x64, 0x74, 0x8B)
COLOR_CARD = RGBColor(0xEF, 0xF4, 0xFF)

LAYOUT_TITLE = 0
LAYOUT_TITLE_ONLY = 5
LAYOUT_BLANK = 6
LAYOUT_CONTENT = 1


@dataclass(slots=True)
class MetricBox:
    label: str
    value: str
    note: str = ""


@dataclass(slots=True)
class SlideTable:
    columns: list[str]
    rows: list[list[str]]


@dataclass(slots=True)
class SlideSpec:
    kind: str
    title: str = ""
    subtitle: str = ""
    bullets: list[str] = field(default_factory=list)
    metrics: list[MetricBox] = field(default_factory=list)
    chart: bytes | None = None
    caption: str = ""
    table: SlideTable | None = None
    notes: str = ""


@dataclass(slots=True)
class DeckSpec:
    title: str
    subtitle: str = ""
    slides: list[SlideSpec] = field(default_factory=list)


# --------------------------------------------------------------------------- #
# Tiện ích vẽ
# --------------------------------------------------------------------------- #
def _set_text(frame, text: str, *, size: int, bold: bool = False,
              color: RGBColor = COLOR_TEXT, align=PP_ALIGN.LEFT) -> None:
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name, run.font.size, run.font.bold = FONT, Pt(size), bold
    run.font.color.rgb = color


def _add_title(slide, text: str, size: int = 30) -> None:
    if slide.shapes.title is not None:
        _set_text(slide.shapes.title.text_frame, text, size=size, bold=True,
                  color=COLOR_TITLE)
        return
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.1), Inches(0.9))
    _set_text(box.text_frame, text, size=size, bold=True, color=COLOR_TITLE)


def _add_footer(slide, text: str, width_in: float) -> None:
    box = slide.shapes.add_textbox(Inches(0.6), Inches(6.85), Inches(width_in - 1.2),
                                   Inches(0.35))
    _set_text(box.text_frame, text, size=10, color=COLOR_MUTED)


# --------------------------------------------------------------------------- #
# Từng kiểu slide
# --------------------------------------------------------------------------- #
def _render_title(presentation, spec: SlideSpec, width_in: float):
    slide = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_TITLE])
    _set_text(slide.shapes.title.text_frame, spec.title, size=40, bold=True,
              color=COLOR_TITLE, align=PP_ALIGN.CENTER)
    if len(slide.placeholders) > 1 and spec.subtitle:
        _set_text(slide.placeholders[1].text_frame, spec.subtitle, size=18,
                  color=COLOR_MUTED, align=PP_ALIGN.CENTER)
    return slide


def _render_summary(presentation, spec: SlideSpec, width_in: float):
    """Các chỉ tiêu chính dạng ô: số lớn, nhãn nhỏ - đọc được từ cuối phòng họp."""
    slide = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_TITLE_ONLY])
    _add_title(slide, spec.title)

    metrics = spec.metrics[:4]
    if not metrics:
        return slide

    margin, gap = 0.7, 0.3
    box_width = (width_in - 2 * margin - gap * (len(metrics) - 1)) / len(metrics)
    for index, metric in enumerate(metrics):
        left = Inches(margin + index * (box_width + gap))
        shape = slide.shapes.add_shape(1, left, Inches(2.0),  # 1 = rectangle
                                       Inches(box_width), Inches(2.0))
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_CARD
        shape.line.color.rgb = COLOR_CARD
        shape.shadow.inherit = False

        frame = shape.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        _set_text(frame, metric.value, size=32, bold=True, color=COLOR_ACCENT,
                  align=PP_ALIGN.CENTER)
        for text, size, color in ((metric.label, 13, COLOR_TEXT),
                                  (metric.note, 11, COLOR_MUTED)):
            if not text:
                continue
            paragraph = frame.add_paragraph()
            paragraph.alignment = PP_ALIGN.CENTER
            run = paragraph.add_run()
            run.text = text
            run.font.name, run.font.size = FONT, Pt(size)
            run.font.color.rgb = color

    if spec.bullets:
        box = slide.shapes.add_textbox(Inches(margin), Inches(4.4),
                                       Inches(width_in - 2 * margin), Inches(2.0))
        _fill_bullets(box.text_frame, spec.bullets, size=14)
    return slide


def _fill_bullets(frame, bullets: list[str], size: int = 16) -> None:
    frame.clear()
    frame.word_wrap = True
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(10)
        run = paragraph.add_run()
        run.text = f"• {bullet}"
        run.font.name, run.font.size = FONT, Pt(size)
        run.font.color.rgb = COLOR_TEXT


def _render_bullet(presentation, spec: SlideSpec, width_in: float):
    slide = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_TITLE_ONLY])
    _add_title(slide, spec.title)
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.7),
                                   Inches(width_in - 1.6), Inches(4.6))
    _fill_bullets(box.text_frame, spec.bullets)
    return slide


def _render_chart(presentation, spec: SlideSpec, width_in: float):
    slide = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_TITLE_ONLY])
    _add_title(slide, spec.title)

    if spec.chart:
        picture_width = Inches(width_in - 2.4)
        slide.shapes.add_picture(io.BytesIO(spec.chart), Inches(1.2), Inches(1.6),
                                 width=picture_width)
    if spec.caption:
        box = slide.shapes.add_textbox(Inches(1.2), Inches(6.3),
                                       Inches(width_in - 2.4), Inches(0.5))
        _set_text(box.text_frame, spec.caption, size=11, color=COLOR_MUTED,
                  align=PP_ALIGN.CENTER)
    return slide


def _render_table(presentation, spec: SlideSpec, width_in: float):
    slide = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_TITLE_ONLY])
    _add_title(slide, spec.title)
    if spec.table is None:
        return slide

    rows = spec.table.rows

    shape = slide.shapes.add_table(
        len(rows) + 1, len(spec.table.columns),
        Inches(0.8), Inches(1.7), Inches(width_in - 1.6),
        Inches(0.4 * (len(rows) + 1)),
    )
    table = shape.table
    for index, column in enumerate(spec.table.columns):
        cell = table.cell(0, index)
        _set_text(cell.text_frame, str(column), size=12, bold=True, color=COLOR_TITLE)
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row[: len(spec.table.columns)]):
            _set_text(table.cell(row_index, column_index).text_frame, str(value), size=12)

    if spec.caption:
        box = slide.shapes.add_textbox(Inches(0.8), Inches(6.6),
                                       Inches(width_in - 1.6), Inches(0.4))
        _set_text(box.text_frame, spec.caption, size=11, color=COLOR_MUTED)
    return slide


# Số dòng dữ liệu vừa một slide. Bảng bắt đầu ở 1.7", mỗi dòng 0.4", chừa chỗ cho
# dòng chú thích ở 6.6" => (6.6 - 1.7) / 0.4 ≈ 12 dòng kể cả tiêu đề cột.
MAX_TABLE_ROWS = 11


def _paginate_tables(specs: list[SlideSpec]) -> list[SlideSpec]:
    """Bảng dài thành nhiều slide thay vì bị cắt cụt.

    Trước đây dôi ra bao nhiêu dòng thì bỏ bấy nhiêu, kèm một câu "xem chi tiết
    trong báo cáo". Câu đó sai ở hai mặt: người xem slide không có bản báo cáo
    trong tay, và một bộ slide tạo riêng lẻ thì không có bản báo cáo nào cả - 25
    trong 33 thiết bị biến mất mà chỗ duy nhất nhắc tới chúng lại trỏ vào hư không.

    Bảng số liệu là thứ phải ĐỦ: thiếu một dòng là người đọc cộng ra số khác với
    tổng ghi ở slide trước đó.
    """
    result: list[SlideSpec] = []
    for spec in specs:
        if spec.kind != "table" or spec.table is None:
            result.append(spec)
            continue

        rows = spec.table.rows
        pages = [rows[i:i + MAX_TABLE_ROWS] for i in range(0, len(rows), MAX_TABLE_ROWS)] or [[]]
        for index, page in enumerate(pages, start=1):
            # Đánh số trang ngay trên tiêu đề: người xem biết còn slide nữa, và
            # biết mình đang ở đâu trong bảng.
            title = spec.title if len(pages) == 1 else f"{spec.title} ({index}/{len(pages)})"
            result.append(replace(
                spec,
                title=title,
                table=SlideTable(columns=spec.table.columns, rows=page),
                # Ghi chú người trình bày chỉ gắn vào slide đầu, không lặp lại.
                notes=spec.notes if index == 1 else "",
                caption=spec.caption if len(pages) == 1 else
                        f"{spec.caption + ' - ' if spec.caption else ''}"
                        f"dòng {(index - 1) * MAX_TABLE_ROWS + 1}-"
                        f"{(index - 1) * MAX_TABLE_ROWS + len(page)}/{len(rows)}",
            ))
    return result


def count_slides(specs: list[SlideSpec]) -> int:
    """Số slide mà `build_pptx` sẽ thật sự dựng ra từ danh sách này.

    Không bằng `len(specs)`: một bảng dài nở ra nhiều slide ở bước phân trang. Ai
    muốn báo con số cho người dùng thì phải hỏi qua đây, nếu không giao diện sẽ
    ghi "5 slide" trên một file 7 slide - và người dùng tin vào cái nhãn chứ không
    mở file ra đếm.
    """
    return len([spec for spec in _paginate_tables(specs) if spec.kind in RENDERERS])


RENDERERS = {
    "title": _render_title,
    "summary": _render_summary,
    "bullet": _render_bullet,
    "chart": _render_chart,
    "table": _render_table,
}


# --------------------------------------------------------------------------- #
def build_pptx(deck: DeckSpec, output_path: str | Path,
               template_path: str | Path | None = None) -> Path:
    """Dựng file .pptx. Có file mẫu thì kế thừa theme của mẫu."""
    output_path = Path(output_path)

    if template_path and Path(template_path).exists():
        presentation = Presentation(str(template_path))
        # Mẫu có thể kèm slide sẵn - xoá để chỉ giữ theme.
        for index in range(len(presentation.slides) - 1, -1, -1):
            xml_slides = presentation.slides._sldIdLst
            xml_slides.remove(list(xml_slides)[index])
    else:
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)      # 16:9
        presentation.slide_height = Inches(7.5)

    width_in = Emu(presentation.slide_width).inches

    for spec in _paginate_tables(deck.slides):
        renderer = RENDERERS.get(spec.kind)
        if renderer is None:
            logger.warning("Bỏ qua slide kiểu lạ: %r", spec.kind)
            continue
        slide = renderer(presentation, spec, width_in)
        if spec.notes:
            slide.notes_slide.notes_text_frame.text = spec.notes

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output_path))
    return output_path
