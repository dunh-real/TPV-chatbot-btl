"""Soạn văn bản giao nhiệm vụ từ bảng phân công của workflow 2.

Ý tưởng: bản soát đã chỉ ra đơn vị nào phải làm gì, hạn nào - mà đó đúng là toàn
bộ phần ruột của một công văn giao việc. Bắt cán bộ đọc bảng trên màn hình rồi gõ
lại vào Word là bắt họ chép tay một thứ hệ thống đã biết.

HOÀN TOÀN TẤT ĐỊNH, không gọi LLM. Phần lời của hai mẫu này là văn khuôn - "nghiêm
túc triển khai thực hiện", "chịu trách nhiệm thi hành Quyết định này" - viết sẵn
đúng hơn và rẻ hơn là hỏi model. Phần thay đổi giữa hai văn bản chỉ là bảng phân
công, mà bảng đó lấy nguyên từ dữ liệu đã có.

Hai mẫu, khác nhau ở thẩm quyền chứ không ở nội dung:
    cong_van     đôn đốc, giao việc thường xuyên trong thẩm quyền sẵn có
    quyet_dinh   giao nhiệm vụ bằng một quyết định hành chính, có căn cứ và điều khoản

Chỗ hệ thống KHÔNG biết thì để dấu chấm lửng đúng lối văn thư ("Số: ...../CV-...")
chứ không đoán. Một số ký hiệu bịa trông y như số thật là thứ nguy hiểm nhất có
thể nhét vào văn bản trình ký.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any, Literal

from app.documents.docx_builder import (
    FONT,
    MARGINS_MM,
    SIZE_PT,
    TITLE_SIZE_PT,
    RenderedTable,
    _apply_spacing,
    append_table,
)

logger = logging.getLogger(__name__)

LoaiVanBan = Literal["cong_van", "quyet_dinh"]

# Chỗ để trống cho người soạn điền. Dấu chấm lửng là quy ước văn thư ai cũng hiểu
# là "điền vào đây", khác hẳn một giá trị bịa trông như thật.
TRONG = "……………………"

COT_BANG = ("TT", "Đơn vị thực hiện", "Nội dung nhiệm vụ", "Số liệu cần chuẩn bị", "Thời hạn")

# Loại văn bản đến -> mẫu nên dùng. Mặc định là công văn: giao việc trong thẩm
# quyền sẵn có thì không cần tới một quyết định hành chính. Chỉ khi văn bản đến
# chính là một quyết định hoặc một đề án/kế hoạch cần phân công chính thức thì
# mẫu quyết định mới đúng tầm.
GOI_Y_QUYET_DINH = {"quyet_dinh", "de_an", "ke_hoach"}


@dataclass(slots=True)
class MetaGiaoViec:
    """Phần thể thức. Trường nào rỗng thì in dấu chấm lửng để người soạn điền."""

    co_quan: str = ""
    co_quan_chu_quan: str = ""
    so_ky_hieu: str = ""
    dia_danh: str = ""
    ngay: str = ""
    trich_yeu: str = ""
    can_cu: list[str] = field(default_factory=list)
    nguoi_ky: str = ""
    chuc_vu_ky: str = ""
    deadline: str | None = None
    mo_dau: str = ""


def goi_y_mau(document_type: str) -> LoaiVanBan:
    """Mẫu nên chọn sẵn cho một loại văn bản đến."""
    return "quyet_dinh" if (document_type or "").strip().lower() in GOI_Y_QUYET_DINH else "cong_van"


def _hoac_trong(value: str) -> str:
    return value.strip() or TRONG


def bang_phan_cong(tasks: list[dict[str, Any]]) -> RenderedTable:
    """Bảng phân công - đúng thứ giao diện hiện và đúng thứ đổ vào file.

    Một nguồn duy nhất cho cả hai nơi: bảng trên màn hình mà lệch bảng trong file
    thì người dùng phải đối chiếu lại từng dòng, và mất luôn cái lợi của việc sinh
    sẵn văn bản.
    """
    rows: list[list[str]] = []
    for stt, task in enumerate(tasks, start=1):
        rows.append([
            str(stt),
            str(task.get("department_name") or task.get("department") or TRONG),
            str(task.get("task") or "").strip() or TRONG,
            ", ".join(str(x) for x in (task.get("data_needed") or [])) or "-",
            str(task.get("deadline") or "") or "-",
        ])
    return RenderedTable(columns=list(COT_BANG), rows=rows)


def _khung(document, meta: MetaGiaoViec, so_ky_hieu: str):
    """Phần đầu dùng chung: cơ quan ban hành, quốc hiệu, số ký hiệu, địa danh.

    Thể thức thật xếp phần này thành hai cột (cơ quan + số ký hiệu bên trái, quốc
    hiệu + địa danh bên phải). Ở đây dựng tuyến tính như `docx_builder` vẫn làm,
    nhưng THỨ TỰ phải đúng: số ký hiệu trước, địa danh và ngày sau. Đảo lại là sai
    thể thức ngay dòng đầu, và đó là chỗ cán bộ văn thư nhìn trước tiên.
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    def para(text, *, bold=False, size=SIZE_PT, align=None, italic=False):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(str(text))
        run.font.name, run.font.size = FONT, Pt(size)
        run.font.bold, run.font.italic = bold, italic
        if align is not None:
            paragraph.alignment = align
        _apply_spacing(paragraph)
        return paragraph

    if meta.co_quan_chu_quan.strip():
        para(meta.co_quan_chu_quan.upper())
    para(_hoac_trong(meta.co_quan).upper(), bold=True)
    para("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para("Độc lập - Tự do - Hạnh phúc", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(f"Số: {so_ky_hieu}")
    ngay = meta.ngay.strip() or date.today().strftime("ngày %d tháng %m năm %Y")
    para(f"{_hoac_trong(meta.dia_danh)}, {ngay}", italic=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    return para


def _chan(para, meta: MetaGiaoViec, nhu: str) -> None:
    """Nơi nhận và chữ ký - phần cuối giống nhau ở cả hai mẫu."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    para("")
    para("Nơi nhận:", bold=True, italic=True)
    para(f"- {nhu};")
    para("- Lưu: VT.")
    para("")
    para(_hoac_trong(meta.chuc_vu_ky).upper(), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para("")
    para("")
    para(_hoac_trong(meta.nguoi_ky), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)


def _noi_nhan(tasks: list[dict[str, Any]]) -> str:
    ten = [str(t.get("department_name") or t.get("department") or "").strip() for t in tasks]
    ten = [t for t in ten if t]
    return "; ".join(ten) if ten else TRONG


def _cong_van(document, meta: MetaGiaoViec, tasks: list[dict[str, Any]]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    para = _khung(document, meta, meta.so_ky_hieu.strip() or "...../CV-" + TRONG)
    para(f"V/v {_hoac_trong(meta.trich_yeu)}", italic=True)
    para("")
    para(f"Kính gửi: {_noi_nhan(tasks)}", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para("")

    if meta.mo_dau.strip():
        para(meta.mo_dau.strip(), align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para(f"{_hoac_trong(meta.co_quan)} đề nghị các đơn vị triển khai thực hiện các "
         f"nhiệm vụ được phân công cụ thể như sau:", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para("")

    append_table(document, bang_phan_cong(tasks))
    para("")

    han = f" trước ngày {meta.deadline}" if meta.deadline else ""
    para(f"Đề nghị các đơn vị nghiêm túc triển khai thực hiện và báo cáo kết quả về "
         f"{_hoac_trong(meta.co_quan)}{han}. Trong quá trình thực hiện, nếu có vướng mắc "
         f"đề nghị phản ánh kịp thời để được hướng dẫn./.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    _chan(para, meta, "Như trên")


def _quyet_dinh(document, meta: MetaGiaoViec, tasks: list[dict[str, Any]]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    para = _khung(document, meta, meta.so_ky_hieu.strip() or "...../QĐ-" + TRONG)
    para("")
    para("QUYẾT ĐỊNH", bold=True, size=TITLE_SIZE_PT, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(f"Về việc giao nhiệm vụ {_hoac_trong(meta.trich_yeu)}",
         bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para("")
    para(_hoac_trong(meta.chuc_vu_ky).upper(), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para("")

    # Căn cứ để trống VẪN in một dòng mẫu: quyết định không có căn cứ là quyết
    # định sai thể thức, để trắng chỗ này thì người soạn rất dễ quên.
    for can_cu in (meta.can_cu or [f"Căn cứ {TRONG}"]):
        para(f"{can_cu.strip().rstrip(';')};", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para(f"Theo đề nghị của {TRONG}.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para("")
    para("QUYẾT ĐỊNH:", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para("")

    para("Điều 1. Giao nhiệm vụ cho các đơn vị có tên dưới đây thực hiện các nội dung "
         "công việc cụ thể như sau:", bold=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    append_table(document, bang_phan_cong(tasks))
    para("")

    han = f" trước ngày {meta.deadline}" if meta.deadline else ""
    para(f"Điều 2. Các đơn vị có tên tại Điều 1 có trách nhiệm tổ chức thực hiện nhiệm vụ "
         f"được giao và báo cáo kết quả về {_hoac_trong(meta.co_quan)}{han}.",
         bold=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para(f"Điều 3. Quyết định này có hiệu lực kể từ ngày ký. Thủ trưởng các đơn vị có tên "
         f"tại Điều 1 và các tổ chức, cá nhân có liên quan chịu trách nhiệm thi hành "
         f"Quyết định này./.", bold=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    _chan(para, meta, "Như Điều 3")


def build_giao_viec(
    loai: LoaiVanBan, meta: MetaGiaoViec, tasks: list[dict[str, Any]], output_path: str | Path
) -> Path:
    """Dựng file .docx theo thể thức Nghị định 30 (phông, cỡ, lề của docx_builder)."""
    import docx
    from docx.shared import Mm, Pt

    if not tasks:
        raise ValueError("Không có nhiệm vụ nào để giao - chưa soạn được văn bản.")

    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name, style.font.size = FONT, Pt(SIZE_PT)
    style.paragraph_format.line_spacing = 1.3

    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = Mm(MARGINS_MM["top"])
    section.bottom_margin = Mm(MARGINS_MM["bottom"])
    section.left_margin = Mm(MARGINS_MM["left"])
    section.right_margin = Mm(MARGINS_MM["right"])

    (_quyet_dinh if loai == "quyet_dinh" else _cong_van)(document, meta, tasks)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path
