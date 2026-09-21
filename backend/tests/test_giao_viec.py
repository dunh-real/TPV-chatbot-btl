"""Soạn công văn / quyết định giao nhiệm vụ từ bảng phân công của workflow 2.

Phần này KHÔNG gọi LLM: lời văn là văn khuôn, số liệu lấy nguyên từ `tasks`. Vì
vậy mọi test ở đây đều là test tất định - chạy bao nhiêu lần cũng ra một kết quả.

Van chặn quan trọng nhất nằm ở `test_khong_bia_so_ky_hieu...`: một số văn bản bịa
trông y như số thật là thứ nguy hiểm nhất có thể nhét vào bản trình ký.
"""

from __future__ import annotations

import httpx
import pytest

from app.documents.giao_viec import (
    TRONG,
    MetaGiaoViec,
    bang_phan_cong,
    build_giao_viec,
    goi_y_mau,
)
from app.main import app

TASKS = [
    {"department": "PHCQT", "department_name": "Phòng Hành chính quản trị",
     "task": "Tổng hợp báo cáo của các đơn vị.",
     "data_needed": ["Danh sách nhân sự", "Biến động tăng/giảm"], "deadline": "20/9/2026"},
    {"department": "PKT", "department_name": "Phòng Kỹ thuật",
     "task": "Thống kê trang thiết bị đang sử dụng.", "data_needed": [], "deadline": None},
]

META = MetaGiaoViec(co_quan="Công ty TPV", dia_danh="Hà Nội",
                    trich_yeu="rà soát nhân sự", chuc_vu_ky="Giám đốc",
                    deadline="25/9/2026")


def _doc(path):
    import docx

    return docx.Document(str(path))


def _than(path) -> list[str]:
    """Các phần tử thân tài liệu theo ĐÚNG thứ tự, bảng hiện thành một mục riêng.

    `.paragraphs` và `.tables` là hai danh sách tách rời nên đọc riêng từng cái
    không nói được bảng nằm ở đâu - mà "bảng nằm ở đâu" chính là thứ cần kiểm.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = _doc(path)
    ra: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            if text := Paragraph(child, document).text.strip():
                ra.append(text)
        elif child.tag.endswith("}tbl"):
            ra.append(f"[BẢNG {len(Table(child, document).rows)} dòng]")
    return ra


# ------------------------------------------------------------ bảng ------- #
def test_bang_lay_dung_du_lieu_cua_tung_nhiem_vu():
    bang = bang_phan_cong(TASKS)

    assert bang.rows[0] == ["1", "Phòng Hành chính quản trị", "Tổng hợp báo cáo của các đơn vị.",
                            "Danh sách nhân sự, Biến động tăng/giảm", "20/9/2026"]
    # Không có số liệu / không có hạn riêng thì in gạch, không để ô trắng: ô trắng
    # trong bảng trình ký trông như bị quên điền.
    assert bang.rows[1][3] == "-" and bang.rows[1][4] == "-"


def test_noi_nhan_ngoai_danh_muc_van_co_ten_trong_bang():
    bang = bang_phan_cong([{"department": "", "department_name": "Ban Giám đốc",
                            "task": "Chỉ đạo thực hiện."}])

    assert bang.rows[0][1] == "Ban Giám đốc"


# ------------------------------------------------------- hai mẫu --------- #
@pytest.mark.parametrize("loai,dau_hieu", [
    ("cong_van", "Kính gửi:"),
    ("quyet_dinh", "QUYẾT ĐỊNH:"),
])
def test_moi_mau_ra_dung_the_thuc_cua_no(tmp_path, loai, dau_hieu):
    than = _than(build_giao_viec(loai, META, TASKS, tmp_path / f"{loai}.docx"))

    assert any(dau_hieu in dong for dong in than)
    assert than[0] == "CÔNG TY TPV"
    # Quốc hiệu và tiêu ngữ là bắt buộc ở mọi văn bản hành chính.
    assert "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" in than
    assert "Độc lập - Tự do - Hạnh phúc" in than


def test_so_ky_hieu_dung_truoc_dia_danh_va_ngay(tmp_path):
    """Đảo hai dòng này là sai thể thức ngay chỗ văn thư nhìn trước tiên."""
    than = _than(build_giao_viec("cong_van", META, TASKS, tmp_path / "cv.docx"))
    so = next(i for i, d in enumerate(than) if d.startswith("Số:"))
    ngay = next(i for i, d in enumerate(than) if d.startswith("Hà Nội,"))

    assert so < ngay


def test_bang_nam_giua_than_bai_khong_roi_xuong_cuoi(tmp_path):
    """Bảng phải đứng ngay sau câu dẫn, trước phần ký - không rơi sau chữ ký."""
    for loai, truoc in (("cong_van", "đề nghị các đơn vị triển khai"),
                        ("quyet_dinh", "Điều 1.")):
        than = _than(build_giao_viec(loai, META, TASKS, tmp_path / f"{loai}.docx"))
        bang = next(i for i, d in enumerate(than) if d.startswith("[BẢNG"))

        assert any(truoc in d for d in than[:bang])
        assert any(d.startswith("Nơi nhận") for d in than[bang:])
        assert "[BẢNG 3 dòng]" == than[bang]      # 1 dòng tiêu đề + 2 nhiệm vụ


def test_quyet_dinh_luon_co_dong_can_cu(tmp_path):
    """Quyết định không căn cứ là sai thể thức; để trắng thì người soạn dễ quên."""
    than = _than(build_giao_viec("quyet_dinh", META, TASKS, tmp_path / "qd.docx"))

    assert any(d.startswith("Căn cứ") for d in than)
    assert [d[:7] for d in than if d.startswith("Điều ")] == ["Điều 1.", "Điều 2.", "Điều 3."]


def test_can_cu_nguoi_dung_dua_thi_dung_cua_nguoi_dung(tmp_path):
    meta = MetaGiaoViec(co_quan="Công ty TPV",
                        can_cu=["Căn cứ Quyết định số 12/QĐ-TPV ngày 01/3/2026"])
    than = _than(build_giao_viec("quyet_dinh", meta, TASKS, tmp_path / "qd.docx"))

    assert "Căn cứ Quyết định số 12/QĐ-TPV ngày 01/3/2026;" in than


# --------------------------------------------- không bịa dữ liệu --------- #
def test_khong_bia_so_ky_hieu_va_nguoi_ky(tmp_path):
    """Trường không biết thì để dấu chấm lửng, tuyệt đối không đặt giá trị trông như thật."""
    than = _than(build_giao_viec("cong_van", MetaGiaoViec(), TASKS, tmp_path / "cv.docx"))
    noi_dung = "\n".join(than)

    assert TRONG in noi_dung
    so = next(d for d in than if d.startswith("Số:"))
    assert TRONG in so or "....." in so
    # Không được có một số hiệu đầy đủ kiểu "Số: 105/CV-BGĐ".
    assert not any(ky_tu.isdigit() for ky_tu in so.replace("/", ""))


def test_khong_co_nhiem_vu_thi_tu_choi_soan(tmp_path):
    with pytest.raises(ValueError, match="Không có nhiệm vụ"):
        build_giao_viec("cong_van", META, [], tmp_path / "rong.docx")


# ------------------------------------------------------- gợi ý mẫu ------- #
@pytest.mark.parametrize("loai_den,mong_doi", [
    ("cong_van_den", "cong_van"), ("thong_bao", "cong_van"), ("bao_cao", "cong_van"),
    ("quyet_dinh", "quyet_dinh"), ("ke_hoach", "quyet_dinh"),
    ("", "cong_van"), ("khong_biet_la_gi", "cong_van"),
])
def test_goi_y_mau_theo_loai_van_ban_den(loai_den, mong_doi):
    assert goi_y_mau(loai_den) == mong_doi


# ---------------------------------------------------------- API --------- #
async def _post(path: str, **kwargs) -> httpx.Response:
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
        return await client.post(path, **kwargs)


async def test_endpoint_tra_ve_duong_tai_ve():
    res = await _post("/api/documents/giao-viec",
                      json={"loai": "quyet_dinh", "co_quan": "Công ty TPV", "tasks": TASKS})

    assert res.status_code == 200, res.text
    body = res.json()
    assert body["loai"] == "quyet_dinh" and body["task_count"] == 2
    assert body["download_url"].startswith("/api/documents/download/")
    assert body["file_name"].startswith("QD_GIAOVIEC")


async def test_endpoint_tu_choi_khi_khong_co_nhiem_vu():
    res = await _post("/api/documents/giao-viec", json={"loai": "cong_van", "tasks": []})

    assert res.status_code == 422
