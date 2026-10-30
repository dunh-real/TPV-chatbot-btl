"""Kiểm tra workflow Word/PowerPoint nhận bảng SQL động, không cần LLM hay GPU."""
from pathlib import Path
from contextlib import asynccontextmanager

import pytest
from docx import Document
from pptx import Presentation
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.agents.nodes import report, drafting, presentation
from app.core.config import get_settings
from app.db.models import Base
from app.schemas.erp import QueryColumn, QueryResult
from app.services.presenton import PresentonError


@pytest.fixture
def query_result():
    return QueryResult(query_id="test", question="Số lượng theo danh mục tháng 9/2026",
        columns=[QueryColumn(name="category", data_type="nvarchar", label="Danh mục"),
                 QueryColumn(name="quantity", data_type="int", label="Số lượng", unit="cái")],
        rows=[["Máy chủ",6],["Máy in",8]], row_count=2, truncated=False,
        schema_version="schema-test", policy_version="policy-test", source_tables=["dbo.Asm_Assets"])


@pytest.fixture
async def environment(tmp_path, monkeypatch, query_result, business_store):
    engine = create_async_engine("sqlite+aiosqlite://")
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    @asynccontextmanager
    async def scope():
        async with factory() as session:
            yield session
            await session.commit()
    questions = []
    async def query(question):
        questions.append(question)
        return query_result
    for module in (report,drafting):
        monkeypatch.setattr(module, "session_scope", scope, raising=False)
        monkeypatch.setattr(module, "query_erp", query)
    monkeypatch.setattr(get_settings(), "output_dir", str(tmp_path))
    yield questions
    await engine.dispose()


def initial():
    return {"request":"Lập báo cáo số lượng theo danh mục tháng 9/2026", "inputs":{"so_ky_hieu":"TEST/BC"},
        "params":{"ky":"2026-09", "thang":9, "nam":2026, "compare_to":"2026-08", "ma_don_vi":[]},
        "ma_don_vi":"", "template":{"ma_template":"BC_ERP", "ten_bao_cao":"Báo cáo ERP", "fields":{}, "file_path":""}}


async def test_report_uses_original_question_and_dynamic_columns(environment):
    state = initial()
    state.update(await report.gather_data_node(state))
    assert environment == [state["request"]]
    assert "queries" in state["data"] and "personnel" not in state["data"]
    state.update(await report.render_node(state))
    state.update(await report.validate_node(state))
    state.update(await report.export_node(state))
    from app.schemas.reports import AggregateResponse
    AggregateResponse(request=state["request"], data=state["data"], sections=state["sections"])
    doc = Document(state["output_path"])
    assert doc.tables[-1].cell(0,0).text == "Danh mục"
    assert doc.tables[-1].cell(1,0).text == "Máy chủ"
    assert doc.tables[-1].cell(1,1).text == "6"


async def test_draft_no_longer_requires_fixed_template_fields(environment):
    state = initial()
    state.update(await drafting.fetch_data_node(state))
    state.update(await drafting.render_sections_node(state))
    state.update(await drafting.validate_node(state))
    state.update(await drafting.export_node(state))
    from app.schemas.reports import DraftResponse
    response = DraftResponse(**{**state, "template": "BC_ERP"})
    assert response.data["queries"][0]["query_id"] == "test"
    assert Path(state["output_path"]).exists()
    assert state["validation"]["status"] == "passed"
    assert state["sections"][0]["table"]["columns"] == ["Danh mục", "Số lượng (cái)"]


async def test_local_powerpoint_contains_every_dynamic_value(environment, monkeypatch):
    async def no_llm(*args, **kwargs):
        return None, {"attempts": []}

    async def local_unit(*args, **kwargs):
        return "Đơn vị thử nghiệm"

    monkeypatch.setattr(presentation, "_summarize", no_llm)
    monkeypatch.setattr(presentation, "_ten_don_vi", local_unit)
    state = initial()
    state.update(await report.gather_data_node(state))
    state.update(await presentation.brief_node(state))
    class OfflinePresenton:
        async def generate(self,*args,**kwargs): raise PresentonError("offline")
        # Đường dựng thẳng (mặc định) cũng phải lùi về bản tự dựng khi Presenton tắt.
        async def get_template(self,*args,**kwargs): raise PresentonError("offline")
    monkeypatch.setattr(presentation,"get_presenton",OfflinePresenton)
    state.update(await presentation.render_node(state))
    state.update(await presentation.verify_node(state))
    from app.schemas.presentations import PresentationResponse
    assert PresentationResponse(**state).data["queries"][0]["query_id"] == "test"
    slides = Presentation(state["output_path"])
    cells = [cell.text for slide in slides.slides for shape in slide.shapes if shape.has_table for row in shape.table.rows for cell in row.cells]
    assert set(["Máy chủ","Máy in","6","8"]) <= set(cells)
    assert state["engine"] == "local" and state["validation"]["status"] == "passed"


@pytest.mark.parametrize("node", [report.gather_data_node,drafting.fetch_data_node])
async def test_documents_reject_partial_query_results(environment, query_result, node):
    query_result.truncated = True
    result = await node(initial())
    assert result.get("error") and not result.get("output_path")
