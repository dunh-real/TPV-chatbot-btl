"""VBQPPL theo phiên bản: các lỗi cấu trúc đã đo trong rag_testing + cây phiên bản + theo dõi sửa đổi.

Không cần GPU/Qdrant/LLM: chunker dùng bộ đếm token xấp xỉ, CSDL là SQLite trong bộ nhớ.
"""

from __future__ import annotations

from datetime import date

import pytest
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.core.config import Settings
from app.db.models import AnswerLog, Base, LegalAction
from app.documents.legal_grammar import parse_legal
from app.legal import service, tracking
from app.legal.amendments import extract_actions, parse_addresses
from app.legal.identity import identify, law_key
from app.legal.query import match_document, parse_question
from app.legal.structure import build_tree, render, scan
from app.rag.chunking import Chunker, approximate_token_counter

# --------------------------------------------------------------------------- #
# Văn bản mẫu (rút gọn từ cấu trúc thật của Luật NHNN 46/2010 và Luật 23/2026)
# --------------------------------------------------------------------------- #
GOC = """| QUỐC HỘI | CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM |
| Luật số: 46/2010/QH12 | *Hà Nội, ngày 16 tháng 6 năm 2010* |

### LUẬT

### NGÂN HÀNG NHÀ NƯỚC VIỆT NAM

### Chương I

### NHỮNG QUY ĐỊNH CHUNG

### Điều 1. Phạm vi điều chỉnh

Luật này quy định về tổ chức và hoạt động của Ngân hàng Nhà nước Việt Nam.

### Điều 4. Nhiệm vụ, quyền hạn của Ngân hàng Nhà nước

1. Xây dựng chiến lược phát triển ngành ngân hàng.

2. Tổ chức thực hiện chính sách tiền tệ quốc gia.

### Chương II

### TÀI CHÍNH

### Điều 5. Vốn pháp định

Vốn pháp định do Thủ tướng Chính phủ quyết định.

### Điều 6. Thu, chi tài chính

1. Thu, chi theo quy định của Chính phủ.

2. Kết quả tài chính được trích lập các quỹ sau đây:

a) Quỹ dự phòng;

b) Quỹ khác.

### Điều 7. Kiểm toán

Kiểm toán nội bộ theo quy định của Thống đốc.

### Điều 8. Hiệu lực thi hành

Luật này có hiệu lực thi hành từ ngày 01 tháng 01 năm 2011.
"""

SUA_DOI = """| QUỐC HỘI | CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM |
| Luật số: 23/2026/QH16 | *Hà Nội, ngày 24 tháng 8 năm 2026* |

### LUẬT

## SỬA ĐỔI, BỔ SUNG MỘT SỐ ĐIỀU CỦA LUẬT NGÂN HÀNG NHÀ NƯỚC VIỆT NAM

### Điều 1. Sửa đổi, bổ sung một số điều của Luật Ngân hàng Nhà nước Việt Nam

1. Sửa đổi, bổ sung khoản 2 Điều 4 như sau:

“2. Tổ chức thực hiện chính sách tiền tệ quốc gia theo quy định của Chính phủ.”.

2. Bổ sung Điều 5a vào sau Điều 5 như sau:

### “Điều 5a. Quỹ dự trữ ngoại hối

Quỹ dự trữ ngoại hối nhà nước do Ngân hàng Nhà nước quản lý.”.

3. Sửa đổi, bổ sung Điều 5 như sau:

“Điều 5. Vốn pháp định

Vốn pháp định do Chính phủ quy định.”.

4. Sửa đổi, bổ sung một số điểm của Điều 6 như sau:

a) Sửa đổi, bổ sung điểm a khoản 2 như sau:

“a) Quỹ dự phòng rủi ro;”;

b) Bổ sung điểm c vào sau điểm b khoản 2 như sau:

“c) Quỹ khen thưởng.”.

5. Thay cụm từ “Thống đốc” bằng cụm từ “Chính phủ” tại Điều 7.

6. Bãi bỏ Điều 7.

### Điều 2. Hiệu lực thi hành

Luật này có hiệu lực thi hành từ ngày 01 tháng 12 năm 2026.
"""


def chunker() -> Chunker:
    return Chunker(Settings(chunk_legal_structure=True), token_counter=approximate_token_counter)


# --------------------------------------------------------------- legal_grammar --- #
def test_chuong_tren_xx_va_ky_tu_eth_duoc_nhan():
    from app.documents.legal_grammar import roman_to_int

    romans = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII", "XIII", "XIV", "XV",
              "XVI", "XVII", "XVIII", "XIX", "XX", "XXI", "XXII"]
    texts: list[str] = []
    for k, r in enumerate(romans, start=1):
        # Bộ luật Hình sự gõ "Ðiều" (U+00D0) ở một chỗ - phải vẫn là Điều.
        texts += [f"Chương {r}", "TÊN CHƯƠNG", f"{'Ðiều' if k == 22 else 'Điều'} {k}. Tên điều", "Nội dung."]
    parsed = parse_legal(texts)
    chuong = [h.text.split()[1] for h in parsed.headings if h.kind == "OWN_CHUONG"]
    assert chuong[-2:] == ["XXI", "XXII"] and len(chuong) == 22
    assert sum(1 for h in parsed.headings if h.kind == "OWN_DIEU") == 22
    assert roman_to_int("XXVI") == 26 and roman_to_int("ABC") == 0


def test_phu_luc_khong_thanh_khoan_cua_dieu_cuoi():
    texts = ["Điều 1. Phạm vi", "1. Khoản một.", "Điều 2. Hiệu lực", "1. Có hiệu lực.",
             "Phụ lục:", "1. Chi phí sản xuất chung", "a) Chi phí nhân viên;"]
    parsed = parse_legal(texts)
    assert [h.kind for h in parsed.headings if h.index >= 4][0] == "ANNEX"
    assert not [h for h in parsed.headings if h.index > 4 and h.kind.startswith("OWN")]


# -------------------------------------------------------------------- chunker --- #
def test_chunker_khong_lam_mat_dong_khoan_mot_dong_va_nhan_dung_dieu():
    chunks = chunker().split(GOC, {})
    noi = "\n".join(c.text for c in chunks)
    for line in ("1. Xây dựng chiến lược phát triển ngành ngân hàng.", "Điều 4. Nhiệm vụ, quyền hạn",
                 "2. Tổ chức thực hiện chính sách tiền tệ quốc gia.", "b) Quỹ khác."):
        assert line in noi
    for c in chunks:
        if "Xây dựng chiến lược" in c.text:
            assert c.section_label.startswith("Điều") and "Điều 4" in " ".join(c.sections)
    assert {u["dieu"] for c in chunks for u in c.units} >= {"1", "4", "5", "6", "7", "8"}


def test_dieu_trich_dan_khong_thanh_muc_cua_luat_sua_doi():
    chunks = chunker().split(SUA_DOI, {})
    sections = [s for c in chunks for s in c.sections]
    assert not any(s.lstrip("“").startswith("Điều 5a") for s in sections)
    assert any(s.startswith("Điều 1.") for s in sections)


def test_heading_rong_khong_bi_bo_o_nhanh_regex():
    text = "### TIÊU ĐỀ MỘT\n\n### TIÊU ĐỀ HAI\n\nNội dung của mục hai."
    chunks = Chunker(Settings(chunk_legal_structure=False), token_counter=approximate_token_counter).split(text)
    assert "TIÊU ĐỀ MỘT" in "\n".join(c.text for c in chunks)


# ------------------------------------------------------------- identity / tree --- #
def test_dinh_danh_van_ban():
    ident = identify(GOC, "d1", "goc.doc")
    assert ident.so_hieu == "46/2010/QH12" and ident.loai == "Luật"
    assert ident.legal_id == "vn:luat:46-2010-qh12"
    assert ident.ngay_ban_hanh == date(2010, 6, 16) and ident.ngay_hieu_luc == date(2011, 1, 1)
    sua = identify(SUA_DOI, "d2", "sua.doc")
    assert sua.ngay_hieu_luc == date(2026, 12, 1) and sua.sua_doi == ["Luật Ngân hàng Nhà nước Việt Nam"]
    assert law_key("Luật Ngân hàng Nhà nước Việt Nam") == law_key(f"{ident.loai} {ident.ten}")
    ky = identify(GOC.replace("có hiệu lực thi hành từ ngày 01 tháng 01 năm 2011",
                              "có hiệu lực thi hành kể từ ngày ký ban hành"), "d3", "x.doc")
    assert ky.ngay_hieu_luc == date(2010, 6, 16)
    # Số hiệu bị tách sang dòng sau (ô bảng "Số:␣␣76 /2026/TT-BXD"), và "Căn cứ ... số 95/2015/QH13"
    # đứng trước phải KHÔNG được lấy làm số hiệu của văn bản.
    tt = identify("| BỘ XÂY DỰNG |  |\n| Số:  76 /2026/TT-BXD | CỘNG HOÀ |\n\n## THÔNG TƯ\n\nBan hành định mức\n\n"
                  "Căn cứ Bộ luật Hàng hải Việt Nam số 95/2015/QH13;\n\n### Điều 1. Phạm vi\n\nNội dung.", "t", "tt.docx")
    assert tt.so_hieu == "76/2026/TT-BXD" and tt.legal_id == "vn:tt:76-2026-tt-bxd"
    khong_ro = identify(GOC.replace("Luật này có hiệu lực thi hành từ ngày 01 tháng 01 năm 2011.", ""), "d4", "y")
    assert khong_ro.ngay_hieu_luc is None            # không đoán ngày


def test_cay_component_va_render_giu_du_chu():
    nodes = build_tree(GOC, "vn:luat:46-2010-qh12")
    ids = {n.component_id for n in nodes}
    assert {"vn:luat:46-2010-qh12:dieu-4:khoan-2", "vn:luat:46-2010-qh12:dieu-6:khoan-2:diem-b",
            "vn:luat:46-2010-qh12:chuong-ii"} <= ids
    out = render(nodes)
    for line in GOC.splitlines():
        if line.strip():
            assert line.strip() in out


# ----------------------------------------------------------------- amendments --- #
def test_doc_lenh_sua_doi():
    acts = extract_actions(scan(SUA_DOI))
    ops = [(a.op, a.address_text) for a in acts]
    assert ("sua_doi", "khoản 2 Điều 4") in ops
    assert ("bo_sung", "Điều 5a") in ops
    assert ("sua_doi", "Điều 5") in ops
    assert ("sua_doi", "điểm a khoản 2 Điều 6") in ops            # địa chỉ ghép từ lệnh chứa
    assert ("bo_sung", "điểm c khoản 2 Điều 6") in ops
    assert ("bai_bo", "Điều 7") in ops
    cum_tu = [a for a in acts if a.op == "thay_cum_tu"]
    assert cum_tu and cum_tu[0].reason                              # luôn cách ly
    assert next(a for a in acts if a.address_text == "Điều 5a").anchor == {"dieu": "5"}


def test_lenh_bai_bo_trong_dieu_khoan_thi_hanh():
    """Luật 23/2026 Điều 4: lệnh bãi bỏ văn bản KHÁC nằm trong Điều khoản thi hành, và câu
    lệnh có cụm "đã được sửa đổi, bổ sung" - vẫn phải là bai_bo với đủ hai địa chỉ."""
    text = ("### Điều 1. Phạm vi\n\nNội dung.\n\n### Điều 2. Điều khoản thi hành\n\n"
            "1. Luật này có hiệu lực thi hành từ ngày 01 tháng 12 năm 2026.\n\n"
            "2. Bãi bỏ Điều 22 và khoản 2 Điều 34 của Pháp lệnh Ngoại hối số 28/2005/PL-UBTVQH11 của Ủy ban "
            "thường vụ Quốc hội đã được sửa đổi, bổ sung một số điều theo Pháp lệnh số 06/2013/UBTVQH13.\n")
    acts = extract_actions(scan(text))
    assert [(a.op, a.address_text, a.target_law) for a in acts] == [
        ("bai_bo", "Điều 22", "Pháp lệnh Ngoại hối"), ("bai_bo", "khoản 2 Điều 34", "Pháp lệnh Ngoại hối")]


def test_dia_chi_nhieu_dich():
    addrs = parse_addresses("Bãi bỏ khoản 3 Điều 7, khoản 2 Điều 22 và Điều 49.")
    assert addrs == [{"khoan": "3", "dieu": "7"}, {"khoan": "2", "dieu": "22"}, {"dieu": "49"}]


# --------------------------------------------------------- service + tracking --- #
@pytest.fixture
async def session():
    engine = create_async_engine("sqlite+aiosqlite://")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    maker = async_sessionmaker(engine, expire_on_commit=False)
    async with maker() as s:
        yield s
    await engine.dispose()


async def _index(s, text, doc_id):
    ident = identify(text, doc_id, f"{doc_id}.doc")
    res = await service.index_document(s, ident, text, doc_id, f"{doc_id}.doc")
    await s.flush()
    return res


async def test_ap_lenh_va_tra_theo_ngay(session):
    goc = await _index(session, GOC, "goc")
    lid = goc.legal_id
    res = await _index(session, SUA_DOI, "sua")
    assert lid in res.touched
    truoc, sau = date(2026, 11, 30), date(2026, 12, 1)
    assert "Thủ tướng" in (await service.lookup(session, lid, {"dieu": "5"}, truoc))["text"]
    assert "do Chính phủ quy định" in (await service.lookup(session, lid, {"dieu": "5"}, sau))["text"]
    assert (await service.lookup(session, lid, {"dieu": "5a"}, truoc))["repealed"]
    assert "Quỹ dự trữ ngoại hối" in (await service.lookup(session, lid, {"dieu": "5a"}, sau))["text"]
    assert (await service.lookup(session, lid, {"dieu": "7"}, sau))["repealed"]
    k2 = await service.lookup(session, lid, {"dieu": "4", "khoan": "2"}, sau)
    assert "theo quy định của Chính phủ" in k2["text"] and k2["text"].startswith("### Điều 4")  # kèm tiêu đề Điều
    diem = await service.lookup(session, lid, {"dieu": "6", "khoan": "2", "diem": "c"}, sau)
    assert "Quỹ khen thưởng" in diem["text"] and "trích lập các quỹ" in diem["text"]           # kèm câu dẫn Khoản
    # Đơn vị con không có trong cây (parser bỏ sót) -> lùi về đơn vị cha thay vì trả rỗng
    lui = await service.lookup(session, lid, {"dieu": "6", "khoan": "1", "diem": "a"}, truoc)
    assert lui and lui["approximate"] and lui["component_id"].endswith("dieu-6:khoan-1")
    snaps = await service.snapshots(session, lid)
    assert [(s.valid_from, s.valid_to) for s in snaps] == [(None, sau), (sau, None)]
    assert "Điều 5a" in snaps[1].text and "Điều 5a" not in snaps[0].text


async def test_thu_tu_nap_nguoc_van_dung(session):
    """Nạp văn bản sửa đổi TRƯỚC văn bản gốc: lệnh chờ đích, áp khi gốc tới."""
    await _index(session, SUA_DOI, "sua")
    from sqlalchemy import select
    pending = (await session.execute(select(LegalAction).where(LegalAction.status == "pending_target"))).scalars().all()
    assert pending
    goc = await _index(session, GOC, "goc")
    assert "do Chính phủ quy định" in (await service.lookup(session, goc.legal_id, {"dieu": "5"}, date(2027, 1, 1)))["text"]


async def test_cau_tra_loi_dung_dieu_bi_sua_bi_danh_dau_loi_thoi(session):
    goc = await _index(session, GOC, "goc")
    lid = goc.legal_id
    ids = {}
    for name, cid in (("dieu5", f"{lid}:dieu-5"), ("khoan2", f"{lid}:dieu-4:khoan-2"),
                      ("dieu4", f"{lid}:dieu-4"), ("dieu1", f"{lid}:dieu-1")):
        ids[name] = await tracking.record_answer(session, question=name, answer="...",
                                                 citations=[{"component_ids": [cid], "version_ids": {}}])
    res = await _index(session, SUA_DOI, "sua")
    await tracking.mark_stale(session, res.changes)
    await session.flush()
    status = {k: (await session.get(AnswerLog, v)).status for k, v in ids.items()}
    assert status == {"dieu5": "stale", "khoan2": "stale", "dieu4": "stale", "dieu1": "fresh"}
    assert (await session.get(AnswerLog, ids["dieu5"])).stale_from == date(2026, 12, 1)


def test_khong_co_bang_chung_vbqppl_thi_khong_ghi():
    assert tracking.evidence_from_citations([{"doc_id": "x"}]) == []
    assert tracking.ancestors("vn:l:1:dieu-4:khoan-2:diem-a") == ["vn:l:1:dieu-4:khoan-2", "vn:l:1:dieu-4"]


# ---------------------------------------------------------------------- query --- #
def test_doc_y_dinh_cau_hoi():
    q = parse_question("Điểm b khoản 2 Điều 6 Luật 46/2010/QH12 quy định gì sau khi sửa đổi?")
    assert q.address == {"diem": "b", "khoan": "2", "dieu": "6"} and q.so_hieu == "46/2010/QH12"
    assert q.after_amendment and not q.before_amendment
    q2 = parse_question("Khoản 3 Điều 17 Nghị định 379 quy định gì vào ngày 15/12/2026?")
    assert q2.so == "379" and q2.loai == "Nghị định" and q2.as_of == date(2026, 12, 15)
    q3 = parse_question("Bộ luật Hình sự quy định tuổi chịu trách nhiệm hình sự thế nào?")
    assert q3.name_hint.startswith("Bộ luật Hình sự") and not q3.has_address


def test_khop_van_ban_theo_so_va_ten():
    class D:
        def __init__(self, so, loai, ten):
            self.so_hieu, self.loai, self.ten = so, loai, ten
    docs = [D("379/2026/NĐ-CP", "Nghị định", "Quy định tiêu chuẩn..."), D("100/2015/QH13", "Bộ luật", "HÌNH SỰ")]
    assert match_document(parse_question("Điều 8 Nghị định 379 quy định gì"), docs)[0].so_hieu == "379/2026/NĐ-CP"
    assert match_document(parse_question("Điều 12 Bộ luật Hình sự"), docs)[0].so_hieu == "100/2015/QH13"


# --------------------------------------------------------------------- parser --- #
def test_parser_khong_chet_voi_jc_end(tmp_path):
    import docx
    from docx.oxml.ns import qn

    from app.documents.parser import parse_document

    doc = docx.Document()
    p = doc.add_paragraph("Điều 1. Phạm vi")
    jc = p._p.get_or_add_pPr().makeelement(qn("w:jc"), {qn("w:val"): "end"})
    p._p.get_or_add_pPr().append(jc)
    path = tmp_path / "jc.docx"
    doc.save(path)
    structure = parse_document(path)
    assert structure.blocks[0].alignment == "RIGHT"


# ------------------------------------------------- số hiệu từ OCR / bản ký số --- #
def test_so_hieu_chu_kirin_cua_ocr():
    text = "CHÍNH PHỦ\nSố: 379/2026/NĐ-CР\nNGHỊ ĐỊNH\nQuy định tiêu chuẩn, định mức sử dụng xe ô tô\nĐiều 1. Phạm vi"
    assert identify(text, "d1", "nd.pdf").so_hieu == "379/2026/NĐ-CP"


def test_so_hieu_de_trong_lay_tu_ten_file_khong_lay_van_ban_vien_dan():
    text = ("BỘ NỘI VỤ\nSố:     /2026/TT-BNV\nTHÔNG TƯ\nQuy định biện pháp thi hành Luật Thi đua, khen thưởng\n"
            "và Nghị định số 368/2026/NĐ-CP ngày 25 tháng 9 năm 2026\nĐiều 1. Phạm vi điều chỉnh")
    assert identify(text, "d1", "25.2026.TT-BNV.pdf").so_hieu == "25/2026/TT-BNV"
    assert identify(text, "d1", "tt_25_2026_bnv.pdf").so_hieu == "25/2026/TT-BNV"
    assert identify(text, "d1", "thong_tu.pdf").so_hieu == ""      # không đoán


def test_so_hieu_khac_loai_bi_loai():
    text = "THÔNG TƯ\nHướng dẫn Nghị định số 368/2026/NĐ-CP\nĐiều 1. Phạm vi"
    assert identify(text, "d1", "x.pdf").so_hieu == ""


def test_tach_moc_cau_truc_trang_pdf_bi_gop_dong():
    from app.rag.converter import tach_moc_cau_truc

    text = ("thưởng. 2 **Điều 2. Quy định chung** 1. Đối với cá nhân: a) <mark>Cá nhân</mark> tại khoản 1 "
            "Điều 5. Nội dung; b) Tập thể.")
    assert tach_moc_cau_truc(text).split("\n") == [
        "thưởng. 2", "**Điều 2. Quy định chung**", "1. Đối với cá nhân:",
        "a) Cá nhân tại khoản 1 Điều 5. Nội dung;", "b) Tập thể."]


# ------------------------------------------------ tích hợp cấu hình R12 (rag_testing) --- #
def test_dia_chi_vien_dan_cheo_khong_phai_don_vi_duoc_hoi():
    q = "Hành vi cố ý gây thương tích, nếu không thuộc trường hợp quy định tại Điều 135, thì mức phạt là bao nhiêu?"
    assert parse_question(q).address_is_reference
    assert parse_question("Người lao động thuê lại có quyền gì ngoài các quy định chung tại Điều 5?").address_is_reference
    assert not parse_question("Khoản 2 Điều 5 Luật Ngân hàng Nhà nước quy định gì?").address_is_reference
    assert not parse_question("Điều 5 quy định gì?").address_is_reference


def _khoan(n: int) -> str:
    return f"{n}. " + " ".join(f"nội dung khoản {n} từ {i}" for i in range(70)) + "."


LUAT_DAI = ("| QUỐC HỘI | CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM |\n| Luật số: 99/2026/QH16 | Hà Nội |\n\n### LUẬT\n\n"
            "THỬ NGHIỆM\n\n### Điều 1. Phạm vi\n\nLuật này quy định thử nghiệm.\n\n### Điều 2. Áp dụng\n\n"
            "Áp dụng cho mọi đối tượng.\n\n### Điều 9. Điều dài\n\n" + "\n\n".join(_khoan(n) for n in (1, 2, 3)))


def test_chunk_theo_khoan_khong_gop_xuyen_dieu():
    cs = Chunker(Settings(chunk_legal_structure=True, legal_chunk_granularity="khoan"),
                 token_counter=approximate_token_counter).split(LUAT_DAI)
    dieu9 = [c for c in cs if any(u.get("dieu") == "9" for u in c.units)]
    assert len(dieu9) == 3 and all(c.text.startswith("Điều 9. Điều dài") for c in dieu9)
    assert all(len({u.get("dieu") for u in c.units if u.get("dieu")}) <= 1 for c in cs)
    # Điều 1 và Điều 2 ngắn vẫn mỗi Điều một chunk (không gộp xuyên Điều)
    assert any(c.text.startswith("Điều 1.") and "Điều 2." not in c.text for c in cs)


def test_van_ban_thuong_giu_nguong_chung():
    text = "\n\n".join(" ".join(f"đoạn {k} chữ {i}" for i in range(120)) for k in range(6))
    a = Chunker(Settings(chunk_legal_structure=True, legal_chunk_granularity="khoan"),
                token_counter=approximate_token_counter).split(text)
    b = Chunker(Settings(chunk_legal_structure=True, legal_chunk_granularity="dieu"),
                token_counter=approximate_token_counter).split(text)
    assert [c.text for c in a] == [c.text for c in b]


def test_docx_dung_tu_khoi_khong_qua_markitdown(tmp_path, monkeypatch):
    import docx

    from app.rag.converter import DocumentConverter

    d = docx.Document()
    d.add_paragraph().add_run("Điều 1. Phạm vi điều chỉnh").bold = True
    d.add_paragraph("1. Khoản một quy định việc thử.")
    d.add_paragraph("a) Điểm a của khoản một;")
    path = tmp_path / "vb.docx"
    d.save(path)
    conv = DocumentConverter(Settings(docx_converter="blocks"))

    def khong_duoc_goi(p):
        raise AssertionError("MarkItDown bị gọi")

    monkeypatch.setattr(conv, "_convert_office", khong_duoc_goi)
    text = conv.convert(path).text
    assert "Điều 1. Phạm vi điều chỉnh" in text and "1. Khoản một" in text and "a) Điểm a" in text


async def test_mo_rong_khoan_ve_ca_dieu(monkeypatch):
    from app.agents.nodes import qa
    from app.rag.retrieval import RetrievedChunk

    sib = [{"doc_id": "d", "dieus": ["9"], "chunk_index": i, "text": f"Điều 9. Khoản {i}"} for i in (1, 2, 3)]

    class Store:
        async def scroll_filtered(self, flt, limit=256):
            return sib

    class Retriever:
        store = Store()

    monkeypatch.setattr(qa, "get_retriever", lambda: Retriever())

    def mk(i, dieus, matched_by="rerank"):
        return RetrievedChunk(point_id=str(i), text=f"Khoản {i}", rerank_score=1 - i / 10, rrf_score=0,
                              payload={"doc_id": "d", "dieus": dieus}, matched_by=matched_by)

    out, n = await qa._mo_rong_ve_dieu([mk(2, ["9"]), mk(3, ["9"]), mk(4, []), mk(5, ["9"], "address")], 0)
    assert n == 1 and [c.point_id for c in out] == ["2", "4", "5"]
    assert out[0].text == "Điều 9. Khoản 1\n\nĐiều 9. Khoản 2\n\nĐiều 9. Khoản 3"
    out, n = await qa._mo_rong_ve_dieu([mk(2, ["9"])], 10)          # Điều dài hơn trần: giữ chunk nhỏ
    assert n == 0 and out[0].text == "Khoản 2"


def test_ngu_canh_ghi_trang_thai_hieu_luc_theo_ngay_tra_cuu():
    from app.agents.nodes.qa import build_context
    from app.rag.retrieval import RetrievedChunk

    c = RetrievedChunk(point_id="1", text="Nội dung sửa đổi", rerank_score=1, rrf_score=0,
                       payload={"doc_title": "Luật 23/2026/QH16", "ngay_hieu_luc": "2026-12-01"})
    ctx, _ = build_context([c], 10_000, date(2026, 10, 10))
    assert "CHƯA CÓ HIỆU LỰC tại ngày 10/10/2026 - có hiệu lực từ 01/12/2026" in ctx
    ctx, _ = build_context([c], 10_000, date(2026, 12, 2))
    assert "CHƯA CÓ HIỆU LỰC" not in ctx


TT_CU = ("| BỘ TÀI CHÍNH | CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM |\n"
         "| Số: 118/2025/TT-BTC | Hà Nội, ngày 15 tháng 12 năm 2025 |\n\n"
         "### THÔNG TƯ\n\nQuy định công tác thi đua, khen thưởng trong ngành Tài chính\n\n"
         "### Điều 1. Phạm vi\n\nThông tư này quy định thi đua trong ngành Tài chính theo cách cũ.\n\n"
         "### Điều 2. Hiệu lực thi hành\n\nThông tư này có hiệu lực thi hành từ ngày 01 tháng 01 năm 2026.\n")


def _tt_moi(so: str, bai_bo: bool) -> str:
    cau = ("\n\n2. Bãi bỏ Thông tư số 118/2025/TT-BTC ngày 15 tháng 12 năm 2025 của Bộ trưởng Bộ Tài chính "
           "hướng dẫn Luật Thi đua, khen thưởng số 06/2022/QH15." if bai_bo else "")
    return (f"| BỘ TÀI CHÍNH | CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM |\n"
            f"| Số: {so} | Hà Nội, ngày 01 tháng 10 năm 2026 |\n\n"
            "### THÔNG TƯ\n\nQuy định công tác thi đua, khen thưởng trong ngành Tài chính\n\n"
            "### Điều 1. Phạm vi\n\nThông tư này quy định thi đua trong ngành Tài chính theo cách mới.\n\n"
            "### Điều 2. Hiệu lực thi hành\n\n1. Thông tư này có hiệu lực thi hành kể từ ngày 01 tháng 11 năm 2026."
            + cau + "\n")


def test_doc_lenh_bai_bo_ca_van_ban():
    from app.legal.amendments import extract_doc_actions
    from app.legal.structure import scan

    specs = extract_doc_actions(scan(_tt_moi("146/2026/TT-BTC", True)))
    assert [(s.op, s.target_law) for s in specs] == [("het_hieu_luc", "118/2025/TT-BTC")]   # không gồm 06/2022/QH15
    specs = extract_doc_actions(scan(LUAT_DAI.replace("Áp dụng cho mọi đối tượng.",
                                                      "Bãi bỏ Điều 5 của Thông tư số 01/2020/TT-BTC.")))
    assert specs == []                                                                       # chỉ một phần văn bản
    het = ("Luật Đất đai số 45/2013/QH13 đã được sửa đổi theo Luật số 35/2018/QH14 hết hiệu lực kể từ ngày "
           "Luật này có hiệu lực.")
    specs = extract_doc_actions(scan(LUAT_DAI.replace("Áp dụng cho mọi đối tượng.", het)))
    assert [s.target_law for s in specs] == ["45/2013/QH13"]


async def test_van_ban_bi_bai_bo_het_hieu_luc_tu_ngay_van_ban_moi(session):
    cu = await _index(session, TT_CU, "tt_cu")
    await _index(session, _tt_moi("146/2026/TT-BTC", True), "tt_moi")
    truoc, sau = date(2026, 10, 31), date(2026, 11, 1)
    assert "cách cũ" in (await service.lookup(session, cu.legal_id, {"dieu": "1"}, truoc))["text"]
    assert (await service.lookup(session, cu.legal_id, {"dieu": "1"}, sau))["repealed"]


async def test_thay_the_khong_ghi_ro_chi_de_xuat_cach_ly(session):
    from sqlalchemy import select

    cu = await _index(session, TT_CU, "tt_cu")
    await _index(session, _tt_moi("150/2026/TT-BTC", False), "tt_moi")
    sau = date(2026, 11, 1)
    assert "cách cũ" in (await service.lookup(session, cu.legal_id, {"dieu": "1"}, sau))["text"]   # chưa ẩn
    act = (await session.execute(select(LegalAction).where(LegalAction.op == "het_hieu_luc"))).scalars().one()
    assert act.status == "quarantined" and act.target_law == "118/2025/TT-BTC"
    assert await service.apply_action(session, act, cu.legal_id)                                # người duyệt
    await session.flush()
    assert (await service.lookup(session, cu.legal_id, {"dieu": "1"}, sau))["repealed"]


# ------------------------------------------------------------ xoá tài liệu --- #
async def test_xoa_van_ban_sua_doi_hoan_tac_sua_doi_tren_van_ban_goc(session):
    from sqlalchemy import select

    from app.db.models import LegalComponent, LegalDocument

    goc = await _index(session, GOC, "goc")
    await _index(session, SUA_DOI, "sua")
    sau = date(2026, 12, 1)
    assert "do Chính phủ quy định" in (await service.lookup(session, goc.legal_id, {"dieu": "5"}, sau))["text"]

    rebuild = await service.delete_document(session, "sua")
    assert rebuild == [goc.legal_id]
    assert "Thủ tướng" in (await service.lookup(session, goc.legal_id, {"dieu": "5"}, sau))["text"]
    assert not (await service.lookup(session, goc.legal_id, {"dieu": "7"}, sau))["repealed"]      # bãi bỏ được gỡ
    assert not (await session.execute(select(LegalComponent).where(
        LegalComponent.component_id == f"{goc.legal_id}:dieu-5a"))).scalars().first()               # Điều bổ sung bị gỡ
    assert not (await session.execute(select(LegalAction))).scalars().all()
    assert (await session.execute(select(LegalDocument).where(LegalDocument.doc_id == "sua"))).scalars().first() is None
    snaps = await service.snapshots(session, goc.legal_id)
    assert [(s.valid_from, s.valid_to) for s in snaps] == [(None, None)]


async def test_xoa_van_ban_goc_lenh_sua_doi_quay_ve_cho_dich(session):
    from sqlalchemy import select

    await _index(session, GOC, "goc")
    await _index(session, SUA_DOI, "sua")
    assert await service.delete_document(session, "goc") == []
    actions = (await session.execute(select(LegalAction))).scalars().all()
    assert actions and all(a.status != "applied" and a.target_legal_id is None for a in actions)
    # Nạp lại văn bản gốc: lệnh chờ được áp lại
    goc = await _index(session, GOC, "goc")
    assert "do Chính phủ quy định" in (await service.lookup(session, goc.legal_id, {"dieu": "5"}, date(2027, 1, 1)))["text"]


# ------------------------------------------------- tên văn bản chung dòng tên loại --- #
def test_ten_van_ban_chung_dong_ten_loai():
    text = ("QUỐC HỘI\nLuật số: 24/2008/QH12\nHà Nội, ngày 13 tháng 11 năm 2008\nLUẬT Quốc tịch Việt Nam\n"
            "Căn cứ Hiến pháp nước Cộng hoà xã hội chủ nghĩa Việt Nam\nChương I\nĐiều 1. Quốc tịch Việt Nam")
    ident = identify(text, "d1", "24.2008.QH12.pdf")
    assert (ident.loai, ident.ten, ident.so_hieu) == ("Luật", "Quốc tịch Việt Nam", "24/2008/QH12")
    assert law_key(f"{ident.loai} {ident.ten}") == law_key("Luật Quốc tịch Việt Nam")


async def test_dieu_cua_van_ban_bi_thay_the_neu_ro_van_ban_thay_the(session):
    from app.db.models import LegalDocument
    from app.legal.hooks import _ly_do_khong_hieu_luc

    cu = await _index(session, TT_CU, "tt_cu")
    await _index(session, _tt_moi("146/2026/TT-BTC", True), "tt_moi")
    doc = await session.get(LegalDocument, cu.legal_id)
    text = await _ly_do_khong_hieu_luc(session, doc, "Điều 1", date(2026, 11, 2))
    assert "HẾT HIỆU LỰC từ ngày 01/11/2026" in text and "146/2026/TT-BTC" in text
    assert "chưa có hiệu lực" not in await _ly_do_khong_hieu_luc(session, doc, "Điều 1", date(2026, 12, 1))
