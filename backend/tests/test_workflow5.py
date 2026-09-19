"""Workflow 5: tạo slide.

Hai tầng: bản dựng file .pptx (dùng cho đường lùi) và đường chính - tóm tắt số
liệu do code dựng, Presenton bày ra slide, rồi đối chiếu lại số trong chính file."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.agents.nodes import presentation as pr
from app.agents.nodes import report as rp
from app.documents.charts import compare_bar_chart
from app.documents.pptx_builder import (
    SLIDE_KINDS,
    DeckSpec,
    MetricBox,
    SlideSpec,
    SlideTable,
    build_pptx,
)


# --------------------------------------------------------------- renderer -- #
@pytest.fixture
def deck() -> DeckSpec:
    chart = compare_bar_chart("Quân số", ["Tổng", "Có mặt"], [144, 133], [140, 130])
    return DeckSpec(
        title="Báo cáo quân số tháng 8/2026", subtitle="Phòng Hành chính quản trị",
        slides=[
            SlideSpec(kind="title", title="Báo cáo quân số tháng 8/2026",
                      subtitle="Phòng Hành chính quản trị"),
            SlideSpec(kind="summary", title="Tổng quan",
                      metrics=[MetricBox("Tổng quân số", "144", "+4 (2.9%)"),
                               MetricBox("Có mặt", "133", "92.4%")],
                      bullets=["3/3 đơn vị đã gửi báo cáo"]),
            SlideSpec(kind="chart", title="Biến động quân số", chart=chart,
                      caption="Nguồn: kiểm kê tháng 8/2026"),
            SlideSpec(kind="table", title="Chi tiết theo đơn vị",
                      table=SlideTable(columns=["Đơn vị", "Quân số"],
                                       rows=[["Đơn vị 1", "48"], ["Đơn vị 2", "65"]])),
            SlideSpec(kind="bullet", title="Đánh giá",
                      bullets=["Quân số tăng 4 người", "Tỷ lệ có mặt 92.4%"],
                      notes="Nhấn mạnh trang bị cần bảo dưỡng"),
        ],
    )


def test_dung_du_slide_va_dung_kieu(deck, tmp_path):
    output = build_pptx(deck, tmp_path / "deck.pptx")
    presentation = Presentation(str(output))

    assert len(presentation.slides) == 5
    assert Emu(presentation.slide_width).inches == pytest.approx(13.333, abs=0.01)  # 16:9


def test_bieu_do_va_bang_duoc_nhung_that(deck, tmp_path):
    output = build_pptx(deck, tmp_path / "deck.pptx")
    presentation = Presentation(str(output))

    chart_slide = presentation.slides[2]
    assert any(shape.shape_type == 13 for shape in chart_slide.shapes)   # 13 = picture

    table_slide = presentation.slides[3]
    table = next(shape.table for shape in table_slide.shapes if shape.has_table)
    assert table.cell(0, 0).text_frame.text == "Đơn vị"
    assert table.cell(1, 1).text_frame.text == "48"


def test_khong_co_shape_nao_tran_ra_ngoai(deck, tmp_path):
    """Slide tràn lề là lỗi hay gặp nhất khi dựng pptx bằng code."""
    output = build_pptx(deck, tmp_path / "deck.pptx")
    presentation = Presentation(str(output))
    width, height = presentation.slide_width, presentation.slide_height

    for index, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            if shape.left is None:
                continue
            assert shape.left >= 0 and shape.top >= 0, f"slide {index}"
            assert shape.left + (shape.width or 0) <= width, f"slide {index}"
            assert shape.top + (shape.height or 0) <= height, f"slide {index}"


def test_ghi_chu_nguoi_trinh_bay_duoc_luu(deck, tmp_path):
    output = build_pptx(deck, tmp_path / "deck.pptx")
    presentation = Presentation(str(output))
    assert "bảo dưỡng" in presentation.slides[4].notes_slide.notes_text_frame.text


def test_bang_qua_dai_tach_slide_chu_khong_cat_bo(tmp_path):
    """Bảng số liệu phải ĐỦ: thiếu một dòng là người đọc cộng ra số khác với tổng.

    Trước đây dôi ra bao nhiêu dòng thì bỏ bấy nhiêu, kèm câu "xem chi tiết trong
    báo cáo" - mà bộ slide tạo riêng lẻ thì không có bản báo cáo nào để xem.
    """
    rows = [[f"Đơn vị {i}", str(i)] for i in range(20)]
    deck = DeckSpec(title="x", slides=[
        SlideSpec(kind="table", title="Bảng dài",
                  table=SlideTable(columns=["Đơn vị", "Quân số"], rows=rows))])

    slides = Presentation(str(build_pptx(deck, tmp_path / "deck.pptx"))).slides
    assert len(slides) == 2                       # 20 dòng / 11 mỗi slide

    du_lieu = []
    for slide in slides:
        table = next(shape.table for shape in slide.shapes if shape.has_table)
        # Bỏ dòng tiêu đề cột, dòng này lặp lại ở mọi trang.
        du_lieu += [[c.text for c in row.cells] for row in list(table.rows)[1:]]

    assert du_lieu == rows                        # đủ cả 20 dòng, đúng thứ tự
    assert len({row.cells[0].text for slide in slides
                for row in [list(next(sh.table for sh in slide.shapes
                                      if sh.has_table).rows)[0]]}) == 1  # mọi trang đều có tiêu đề cột


def test_bang_nhieu_trang_co_danh_so_de_biet_dang_o_dau(tmp_path):
    rows = [[f"Đơn vị {i}", str(i)] for i in range(20)]
    deck = DeckSpec(title="x", slides=[
        SlideSpec(kind="table", title="Chi tiết",
                  table=SlideTable(columns=["Đơn vị", "Quân số"], rows=rows))])

    slides = Presentation(str(build_pptx(deck, tmp_path / "deck.pptx"))).slides
    chu = [" ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
           for slide in slides]

    assert "Chi tiết (1/2)" in chu[0] and "Chi tiết (2/2)" in chu[1]
    assert "dòng 1-11/20" in chu[0] and "dòng 12-20/20" in chu[1]


def test_bang_ngan_khong_bi_danh_so_thua(tmp_path):
    deck = DeckSpec(title="x", slides=[
        SlideSpec(kind="table", title="Chi tiết",
                  table=SlideTable(columns=["Đơn vị"], rows=[["A"], ["B"]]))])

    slides = Presentation(str(build_pptx(deck, tmp_path / "deck.pptx"))).slides
    assert len(slides) == 1
    chu = " ".join(sh.text_frame.text for sh in slides[0].shapes if sh.has_text_frame)
    assert "Chi tiết" in chu and "(1/1)" not in chu and "dòng" not in chu


def test_slide_kieu_la_bi_bo_qua(tmp_path):
    deck = DeckSpec(title="x", slides=[
        SlideSpec(kind="title", title="Bìa"),
        SlideSpec(kind="kieu_khong_ton_tai", title="Lạ"),
    ])
    assert len(Presentation(str(build_pptx(deck, tmp_path / "d.pptx"))).slides) == 1



# --------------------------------------------------------------------------- #
# Bản tóm tắt số liệu gửi cho Presenton
#
# Đây là ranh giới của cả workflow: Presenton chỉ nhìn thấy chuỗi này. Thiếu một
# con số ở đây thì slide không thể có nó; thừa một câu ở đây thì model được mời
# suy diễn. Nên phần này kiểm kỹ hơn phần dựng file.
# --------------------------------------------------------------------------- #
DATA = {
    "personnel": {
        "metrics": {
            "total_personnel": {"value": 28, "prev": 27, "delta": 1,
                                "delta_pct": 3.7, "share_pct": None},
            "new_hires": {"value": 1, "prev": None, "delta": None,
                          "delta_pct": None, "share_pct": 3.6},
        },
        "scope": {"units_requested": 2, "units_with_data": 2, "units_missing": [],
                  "nguon": "Hrm_EmployeeProfile", "ghi_chu": "Số liệu kỳ được dựng lại.",
                  "khong_co_chi_tieu": ["có mặt", "vắng"]},
        "breakdown": [
            {"ma_don_vi": "00001", "ten_don_vi": "Đơn vị 1", "quan_so": 11,
             "quan_so_ky_truoc": 10, "tuyen_moi": 1, "nghi_viec": 0},
            {"ma_don_vi": "00002", "ten_don_vi": "Đơn vị 2", "quan_so": 17,
             "quan_so_ky_truoc": 17, "tuyen_moi": 0, "nghi_viec": 0},
        ],
        "consistency": [],
    },
    "reporting": {"units_total": 2, "units_reported": 1,
                  "missing": [{"ma_don_vi": "00001", "ten_don_vi": "Đơn vị 1"}],
                  "reported": []},
}
PARAMS = {"ky": "2026-08", "thang": 8, "nam": 2026, "compare_to": "2026-07"}


def test_tom_tat_co_du_moi_con_so_se_len_slide():
    """Model không được tính gì: chênh lệch và tỷ lệ đều phải có sẵn trong tóm tắt."""
    brief, _ = pr.build_brief(DATA, PARAMS)

    assert "Tổng quân số: 28" in brief
    assert "kỳ trước 27" in brief
    assert "+1 (+3.7%)" in brief
    assert "chiếm 3.6%" in brief
    assert "tháng 8/2026" in brief and "tháng 7/2026" in brief


def test_tom_tat_liet_ke_du_so_lieu_tung_don_vi():
    """Model phải có số của từng đơn vị mới viết nhận xét được."""
    brief, _ = pr.build_brief(DATA, PARAMS)

    assert "Đơn vị 1 / 11 / 10 / 1 / 0" in brief
    assert "Đơn vị 2 / 17 / 17 / 0 / 0" in brief


def test_tom_tat_dan_khong_dung_bang():
    """Template nào của Presenton cũng chặn bảng quá 3-6 dòng, và khi schema từ
    chối thì cả bộ slide về rỗng. Bảng thật do code ghép vào sau."""
    brief, _ = pr.build_brief(DATA, PARAMS)
    assert "KHÔNG dựng thành bảng" in brief
    assert "KHÔNG dựng slide dạng bảng" in pr.INSTRUCTIONS


def test_bang_qua_dai_thi_noi_thang_la_da_cat():
    """Không cắt im lặng: slide phải biết mình đang thiếu dòng."""
    data = {"personnel": {
        "metrics": {}, "scope": {}, "consistency": [],
        "breakdown": [{"ma_don_vi": f"{i:05d}", "ten_don_vi": f"Đơn vị {i}", "quan_so": i,
                       "quan_so_ky_truoc": i, "tuyen_moi": 0, "nghi_viec": 0}
                      for i in range(pr.MAX_BRIEF_ROWS + 5)]}}
    brief, _ = pr.build_brief(data, PARAMS)

    assert "còn 5 dòng nữa không liệt kê" in brief


def test_tom_tat_noi_ro_chi_tieu_khong_co_nguon():
    """Chỉ tiêu không có số liệu phải được nêu, nếu không model sẽ nhận xét về nó."""
    brief, _ = pr.build_brief(DATA, PARAMS)
    assert "Không có số liệu về: có mặt, vắng." in brief


def test_tom_tat_khong_mang_theo_cau_hoi_goc():
    """Câu người dùng gõ ("cho đẹp vào", "chi tiết hơn") là lời mời thêm thắt."""
    brief, _ = pr.build_brief(DATA, PARAMS)
    assert "slide" not in brief.lower() or "đề xuất" not in brief.lower()
    assert "Ghi chú: Số liệu kỳ được dựng lại." in brief


def test_canh_bao_du_lieu_thanh_muc_rieng():
    data = {**DATA, "personnel": {**DATA["personnel"],
                                  "consistency": [{"ma_don_vi": "*", "message": "Quân số lệch."}]}}
    brief, _ = pr.build_brief(data, PARAMS)
    assert "SỐ LIỆU CẦN KIỂM TRA LẠI" in brief and "Quân số lệch." in brief


def test_so_slide_de_xuat_nam_trong_khoang():
    assert pr.MIN_SLIDES <= pr.build_brief(DATA, PARAMS)[1] <= pr.MAX_SLIDES
    assert pr.MIN_SLIDES <= pr.build_brief({}, PARAMS)[1] <= pr.MAX_SLIDES


# --------------------------------------------------------------------------- #
# Đối chiếu số TRÊN FILE đã dựng
# --------------------------------------------------------------------------- #
def _lam_file(tmp_path, *bullets: str) -> str:
    deck = DeckSpec(title="Báo cáo", slides=[
        SlideSpec(kind="title", title="Báo cáo quân số tháng 8/2026"),
        SlideSpec(kind="bullet", title="Đánh giá", bullets=list(bullets)),
    ])
    return str(build_pptx(deck, tmp_path / "deck.pptx"))


async def test_so_bia_tren_slide_bi_neu_ten(tmp_path):
    """Presenton giữ nguyên chữ nó viết, nên van chắn phải NÓI RA chứ không xoá."""
    path = _lam_file(tmp_path, "Tổng quân số 28 người", "Có 47 đơn vị chưa gửi báo cáo")

    result = await pr.verify_node({"data": DATA, "params": PARAMS, "output_path": path,
                                   "engine": "presenton", "verify_upto": 2})

    assert result["validation"]["status"] == "warning"
    assert result["slide_count"] == 2
    [issue] = result["validation"]["issues"]
    assert issue["numbers"] == ["47"]
    assert "47" in issue["quote"]


async def test_slide_toan_so_that_thi_khong_canh_bao(tmp_path):
    path = _lam_file(tmp_path, "Tổng quân số 28 người, tăng 1 so với kỳ trước")

    result = await pr.verify_node({"data": DATA, "params": PARAMS, "output_path": path,
                                   "engine": "presenton", "verify_upto": 2})
    assert result["validation"]["status"] == "passed"


async def test_so_trong_bang_cung_duoc_soi(tmp_path):
    """Bảng là chỗ dễ lọt nhất: nó trông như dữ liệu nên không ai đọc kỹ."""
    deck = DeckSpec(title="x", slides=[
        SlideSpec(kind="table", title="Chi tiết",
                  table=SlideTable(columns=["Đơn vị", "Quân số"],
                                   rows=[["Đơn vị 1", "11"], ["Đơn vị 2", "99"]]))])
    path = str(build_pptx(deck, tmp_path / "d.pptx"))

    result = await pr.verify_node({"data": DATA, "params": PARAMS, "output_path": path,
                                   "engine": "presenton", "verify_upto": 1})
    assert ["99"] in [i["numbers"] for i in result["validation"]["issues"]]


async def test_khong_doc_lai_duoc_file_thi_noi_la_chua_kiem(tmp_path):
    """Im lặng ở đây nguy hiểm hơn: người dùng tưởng đã đối chiếu xong."""
    hong = tmp_path / "hong.pptx"
    hong.write_bytes(b"khong phai pptx")

    result = await pr.verify_node({"data": DATA, "params": PARAMS, "verify_upto": 2,
                                   "output_path": str(hong), "engine": "presenton"})
    assert result["validation"]["status"] == "skipped"
    assert result["validation"]["ghi_chu"]


# --------------------------------------------------------------------------- #
# Toàn bộ workflow 5, với Presenton giả
# --------------------------------------------------------------------------- #
@pytest.fixture
async def ppt_env(tmp_path, monkeypatch, erp_session):
    """Số liệu đọc từ ERP tạm, sổ văn bản ở CSDL app tạm, output vào thư mục tạm."""
    from contextlib import asynccontextmanager

    from app.db.models import Base

    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    app_session = factory()

    @asynccontextmanager
    async def _scope():
        yield app_session

    @asynccontextmanager
    async def _erp_scope():
        yield erp_session

    class _Settings:
        output_dir = str(tmp_path)
        utility_model = "test"
        presenton_n_slides = 0

    import app.db.session as db_session

    monkeypatch.setattr(db_session, "session_scope", _scope)
    monkeypatch.setattr(rp, "session_scope", _scope)
    monkeypatch.setattr(rp, "erp_session_scope", _erp_scope)
    monkeypatch.setattr(rp, "get_settings", lambda: _Settings())
    monkeypatch.setattr(pr, "get_settings", lambda: _Settings())
    yield tmp_path
    await app_session.close()
    await engine.dispose()


class ParamsLLM:
    """Chỉ còn MỘT lời gọi LLM trong workflow 5: bước trích tham số."""

    def __init__(self) -> None:
        self.calls = 0

    async def chat_json(self, messages, **kwargs):
        self.calls += 1
        return {"thang": 8, "nam": 2026, "ma_don_vi": [], "so_sanh_thang": None,
                "noi_dung": ["quan_so", "trang_bi"]}


class FakePresenton:
    """Presenton giả: nhận tóm tắt, trả về một file .pptx thật để còn đọc lại."""

    def __init__(self, tmp_path, bullets: list[str] | None = None,
                 loi: Exception | None = None) -> None:
        self.tmp_path = tmp_path
        self.bullets = bullets or []
        self.loi = loi
        self.brief = ""
        self.instructions = ""
        self.n_slides = 0

    async def generate(self, content, *, n_slides, instructions="", **kwargs):
        from app.services.presenton import Deck

        self.brief, self.instructions, self.n_slides = content, instructions, n_slides
        if self.loi:
            raise self.loi
        deck = DeckSpec(title="Báo cáo", slides=[
            SlideSpec(kind="title", title="Báo cáo quân số tháng 8/2026"),
            SlideSpec(kind="bullet", title="Đánh giá", bullets=self.bullets),
        ])
        path = build_pptx(deck, self.tmp_path / "presenton.pptx")
        return Deck(presentation_id="p1", content=Path(path).read_bytes(),
                    remote_path="/app_data/exports/p1.pptx", edit_path="/presentation?id=p1",
                    elapsed_seconds=1.0)


async def test_tao_bo_slide_qua_presenton(ppt_env, monkeypatch):
    from app.agents.graph import run_presentation_workflow

    llm = ParamsLLM()
    gia = FakePresenton(ppt_env, ["Tổng quân số 5 người, tăng 1 so với kỳ trước"])
    monkeypatch.setattr(rp, "get_llm", lambda: llm)
    monkeypatch.setattr(pr, "get_presenton", lambda: gia)

    result = await run_presentation_workflow("Tạo slide báo cáo quân số tháng 8/2026")

    assert result["error"] == ""
    assert result["engine"] == "presenton"
    assert result["validation"]["status"] == "passed"
    assert Path(result["output_path"]).exists()
    assert result["edit_url"] == "/presentation?id=p1"

    # 2 slide của Presenton + bảng chi tiết quân số do code ghép vào.
    assert result["slide_count"] == 3
    presentation = Presentation(result["output_path"])
    bang = [sh.table for s in presentation.slides for sh in s.shapes if sh.has_table]
    assert len(bang) == 1
    assert len(bang[0].rows) == 1 + 2          # tiêu đề cột + 2 đơn vị, không thiếu dòng

    # Đúng một lời gọi LLM nội bộ: bước trích tham số. Phần chữ trên slide không
    # còn đi qua model của hệ thống nữa.
    assert llm.calls == 1
    # Và Presenton chỉ thấy số liệu, không thấy câu hỏi gốc.
    assert "Tạo slide báo cáo" not in gia.brief
    assert "Tổng quân số" in gia.brief
    assert "KHÔNG tự cộng" in gia.instructions


async def test_presenton_hong_thi_van_ra_file_bang_duong_lui(ppt_env, monkeypatch):
    """Hỏng ở bên kia không được làm người dùng ra về tay trắng."""
    from app.agents.graph import run_presentation_workflow
    from app.services.presenton import PresentonError

    llm = ParamsLLM()
    gia = FakePresenton(ppt_env, loi=PresentonError("container chưa chạy"))
    monkeypatch.setattr(rp, "get_llm", lambda: llm)
    monkeypatch.setattr(pr, "get_presenton", lambda: gia)

    result = await run_presentation_workflow("Tạo slide báo cáo quân số tháng 8/2026")

    assert result["engine"] == "local"
    assert Path(result["output_path"]).exists()
    assert result["slide_count"] > 1
    # Nói rõ bộ slide này khác thứ người dùng chờ đợi, và vì sao.
    assert any("container chưa chạy" in a for a in result["assumptions"])
    # Đường lùi không gọi LLM viết chữ: vẫn đúng một lời gọi của bước tham số.
    assert llm.calls == 1


async def test_duong_lui_van_du_so_lieu(ppt_env, monkeypatch):
    """Bộ slide đường lùi khô khan nhưng không được thiếu bảng chi tiết."""
    from app.agents.graph import run_presentation_workflow
    from app.services.presenton import PresentonError

    monkeypatch.setattr(rp, "get_llm", lambda: ParamsLLM())
    monkeypatch.setattr(pr, "get_presenton",
                        lambda: FakePresenton(ppt_env, loi=PresentonError("hỏng")))

    result = await run_presentation_workflow("Tạo slide báo cáo tháng 8/2026")
    presentation = Presentation(result["output_path"])

    assert any(sh.has_table for slide in presentation.slides for sh in slide.shapes)
    chu = " ".join(sh.text_frame.text for slide in presentation.slides
                   for sh in slide.shapes if sh.has_text_frame)
    assert "Đơn vị 1" in chu


# --------------------------------------------------------------------------- #
# Đường lùi dựng bằng code: không một chữ nào của LLM
# --------------------------------------------------------------------------- #
def test_duong_lui_khong_co_cau_nao_do_model_viet():
    deck = pr._fallback_deck(DATA, PARAMS)
    bullets = [b for s in deck.slides for b in s.bullets]

    # Mọi câu trong bộ slide đường lùi đều phải truy được về dữ liệu gốc.
    assert bullets
    assert all("Đơn vị" in b or "đơn vị" in b for b in bullets)
    assert [box.value for s in deck.slides for box in s.metrics] == ["28", "1"]


def test_o_chi_tieu_giu_nguyen_bien_dong_da_tinh_san():
    box = pr._metric_box("total_personnel", DATA["personnel"]["metrics"]["total_personnel"])
    assert box.value == "28" and box.note == "+1 (3.7%)"


def test_bang_rong_khong_thanh_slide_bang():
    """Kỳ chưa có trang bị nào: bảng rỗng thì slide chỉ còn dòng tiêu đề cột."""
    assert pr._slide_table("equipment_breakdown", {"equipment": {"breakdown": []}}) is None
    assert pr._slide_table("personnel_breakdown", DATA) is not None


# --------------------------------------------------------------------------- #
# Đếm slide: nhãn trên giao diện phải khớp file
# --------------------------------------------------------------------------- #
def test_slide_count_dem_ca_slide_sinh_them_khi_phan_trang():
    """`len(specs)` đếm trước bước phân trang, nên bảng 33 dòng làm nhãn sai."""
    from app.documents.pptx_builder import count_slides

    specs = [
        SlideSpec(kind="title", title="Bìa"),
        SlideSpec(kind="table", title="Chi tiết",
                  table=SlideTable(columns=["A"], rows=[[str(i)] for i in range(33)])),
    ]
    assert count_slides(specs) == 1 + 3          # 33 dòng / 11 = 3 slide bảng


def test_slide_count_khop_voi_file_dung_ra(tmp_path):
    from app.documents.pptx_builder import count_slides

    specs = [
        SlideSpec(kind="title", title="Bìa"),
        SlideSpec(kind="table", title="Chi tiết",
                  table=SlideTable(columns=["A"], rows=[[str(i)] for i in range(33)])),
        SlideSpec(kind="kieu_la", title="Bị bỏ qua"),
    ]
    output = build_pptx(DeckSpec(title="x", slides=specs), tmp_path / "d.pptx")
    assert count_slides(specs) == len(Presentation(str(output)).slides)


def test_ghep_bang_dai_vao_file_co_san_khong_mat_dong(tmp_path):
    """Bảng 33 dòng ghép vào bộ slide Presenton phải còn đủ 33 dòng.

    Đây là lý do bảng không giao cho Presenton: schema của nó chặn ở 6 dòng, còn
    ở đây bảng dài chỉ nở thêm trang chứ không mất dòng nào.
    """
    from app.documents.pptx_builder import append_to_pptx

    goc = build_pptx(DeckSpec(title="x", slides=[SlideSpec(kind="title", title="Bìa")]),
                     tmp_path / "goc.pptx")
    rows = [[f"Đơn vị {i}", str(i)] for i in range(33)]
    them = append_to_pptx(goc, [SlideSpec(kind="table", title="Chi tiết",
                                          table=SlideTable(columns=["Đơn vị", "Quân số"],
                                                           rows=rows))])

    presentation = Presentation(str(goc))
    assert them == 3                                  # 33 dòng tách thành 3 trang
    assert len(presentation.slides) == 1 + 3          # slide cũ được giữ nguyên
    dem = sum(len(sh.table.rows) - 1 for s in presentation.slides
              for sh in s.shapes if sh.has_table)
    assert dem == 33


def test_chi_dan_nhac_muc_canh_bao_khi_that_su_co_canh_bao():
    """Dặn cứng thì kỳ nào sạch số liệu cũng mọc dòng "cần kiểm tra lại: chưa nêu"."""
    assert "kiểm tra lại" not in pr.instructions_for(DATA)

    co_canh_bao = {**DATA, "personnel": {**DATA["personnel"],
                                         "consistency": [{"ma_don_vi": "*",
                                                          "message": "Quân số lệch."}]}}
    assert "kiểm tra lại" in pr.instructions_for(co_canh_bao)


async def test_bang_ghep_them_khong_bi_soi_thanh_so_bia(tmp_path):
    """Chú thích phân trang do CODE viết, không phải model - soi nó là báo oan.

    Tầng live đỏ đúng chỗ này: bảng 33 dòng tách ba trang, mỗi trang mang chú
    thích "Nguồn: Asm_Assets - dòng 12-22/33", và van chắn đọc 12, 22, 33 thành
    ba con số không truy được về dữ liệu gốc.
    """
    from app.documents.pptx_builder import append_to_pptx

    path = _lam_file(tmp_path, "Tổng quân số 28 người")
    append_to_pptx(path, [SlideSpec(
        kind="table", title="Chi tiết", caption="Nguồn: Asm_Assets",
        table=SlideTable(columns=["Đơn vị", "Số lượng"],
                         rows=[[f"Đơn vị {i}", str(i)] for i in range(33)]))])

    result = await pr.verify_node({"data": DATA, "params": PARAMS, "output_path": path,
                                   "engine": "presenton", "verify_upto": 2})

    assert result["validation"]["status"] == "passed"
    assert result["slide_count"] == 2 + 3        # vẫn đếm đủ slide trong file
