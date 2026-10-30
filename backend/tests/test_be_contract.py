"""Kiểm thử hợp đồng BE: không nhân bản, không ghi đè tác vụ, giữ nội dung DOCX."""
import io
from contextlib import asynccontextmanager
from datetime import datetime, timezone, timedelta
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from fastapi import BackgroundTasks, UploadFile
from sqlalchemy import Column, Integer, String, Boolean, DateTime, MetaData, Table, create_engine, select

from app.core.context import Principal, use_principal
from app.db.erp_auth import AccountAccess
from app.db.ai_repository import AIRepository, now, lease_now
from app.db import ai_session


def who(tenant=94, user=407):
    return Principal(tenant_id=tenant,user_id=str(user),access=AccountAccess(user,tenant,'test','test','',
        permissions=frozenset({'Ai.ReportSummary.Create','Ai.AiChatbot'})))


def repository_patch(monkeypatch, module, repo):
    @asynccontextmanager
    async def repository(): yield repo
    monkeypatch.setattr(module,'repository',repository)


async def test_attachment_without_ingest(monkeypatch,tmp_path):
    from app.api import agent
    path=tmp_path/'source.txt';path.write_text('Nội dung đính kèm')
    ref=SimpleNamespace(path=path,as_dict=lambda:dict(file_id='upload:source.txt',name='source.txt',kind='upload',size=10))
    monkeypatch.setattr(agent.storage,'save_upload',lambda *a,**k:ref)
    background=BackgroundTasks()
    await agent.upload(background,UploadFile(filename='source.txt',file=io.BytesIO(b'test')),kind='upload',ingest=False)
    assert not background.tasks


async def test_same_erp_document_reindexes_without_inserting(monkeypatch):
    from app.services import ai_documents
    from app.rag.ingestion import IngestResult
    monkeypatch.setattr(ai_documents,'enabled',lambda:True)
    repo=SimpleNamespace(lock=AsyncMock(),erp_document=AsyncMock(return_value={'AiServerDocId':'hash422'}),
        patch=AsyncMock(),principal=who(),add=AsyncMock())
    repository_patch(monkeypatch,ai_documents,repo)
    writes=[]
    @ai_documents.tracked_ingestion
    async def ingest(self,text,title,doc_id,source,doc_type,metadata,**kwargs):
        writes.append(doc_id);return IngestResult(doc_id,title,3,1)
    token=ai_session.erp_document.set((422,1))
    try:
        await ingest(None,'Điều 1 nội dung','Tên',source='source.docx')
        await ingest(None,'Điều 1 nội dung','Tên',source='source.docx')
    finally:ai_session.erp_document.reset(token)
    assert writes==['hash422','hash422']
    repo.add.assert_not_awaited()
    assert all(set(c.kwargs)<= {'AiServerDocId','IndexStatus','ProcessingStatus','ChunkCount','LastIndexedTime'} for c in repo.patch.call_args_list)


async def test_sql_and_old_ids_resolve_before_filter(monkeypatch):
    from app.services import ai_documents
    from qdrant_client import models
    monkeypatch.setattr(ai_documents,'enabled',lambda:True)
    async def document(key): return {'AiServerDocId':'hash422','Id':422} if key in {'sql:422','hash422'} else None
    repo=SimpleNamespace(document_ids=AsyncMock(return_value=['hash422']),document=document)
    repository_patch(monkeypatch,ai_documents,repo)
    result=await ai_documents.restrict(models.Filter(must=[models.FieldCondition(key='doc_id',match=models.MatchAny(any=['sql:422','sql:999']))]))
    assert result.must[0].must[0].match.any==['hash422']
    assert await ai_documents.public_ids(['hash422'])=={'hash422':'sql:422'}


async def test_delete_only_index(monkeypatch):
    from app.api import documents
    monkeypatch.setattr(ai_session,'enabled',lambda:True)
    repo=SimpleNamespace(document=AsyncMock(return_value={'Id':422,'AiServerDocId':'hash422'}),lock=AsyncMock(),patch=AsyncMock())
    repository_patch(monkeypatch,ai_session,repo)
    store=SimpleNamespace(delete_document=AsyncMock())
    monkeypatch.setattr(documents,'get_vector_store',lambda:store)
    assert await documents.remove_document('sql:422')==[]
    store.delete_document.assert_awaited_once_with('hash422')
    assert repo.patch.call_args.kwargs==dict(IndexStatus='NotIndexed',ChunkCount=0)


def test_db_clock_and_worker_clock_are_separate():
    utc=datetime.now(timezone.utc).replace(tzinfo=None)
    assert abs((now()-utc-timedelta(hours=7)).total_seconds())<2
    assert abs((lease_now()-utc).total_seconds())<2


def test_docx_content_controls_keep_articles(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    from app.rag.converter import docx_blocks_markdown
    from app.rag.chunking import Chunker,approximate_token_counter
    doc=Document();doc.add_heading('QUY CHẾ LÀM VIỆC',1)
    paras=[doc.add_paragraph('Điều 1. Thời gian làm việc'),doc.add_paragraph('Buổi sáng từ 07 giờ 30 đến 11 giờ 30.'),
           doc.add_paragraph('Điều 2. Trách nhiệm'),doc.add_paragraph('Cán bộ thực hiện đúng nội quy.')]
    control=OxmlElement('w:sdt');content=OxmlElement('w:sdtContent');control.append(content)
    for p in paras:content.append(p._p)
    doc._element.body.insert(1,control)
    path=tmp_path/'source.docx';doc.save(path)
    text=docx_blocks_markdown(path)
    chunks=Chunker(token_counter=approximate_token_counter).split(text)
    combined='\n'.join(c.text for c in chunks)
    assert '07 giờ 30' in combined and 'Điều 2' in combined and 'đúng nội quy' in combined


@pytest.mark.parametrize('fails',[False,True])
async def test_erp_task_preserves_be_fields(monkeypatch,fails):
    from app.services import erp_tasks,ai_outputs
    metadata=MetaData()
    task=Table('AI_Tasks',metadata,Column('Id',Integer,primary_key=True),Column('TenantId',Integer),
        Column('UserId',Integer),Column('IsDeleted',Boolean),Column('Status',String),Column('ResultJson',String),
        Column('ProgressPercentage',Integer),Column('RelatedEntityId',Integer),Column('NotifyOnFinish',Boolean))
    engine=create_engine('sqlite://');metadata.create_all(engine)
    with engine.begin() as conn:
        conn.execute(task.insert().values(Id=554,TenantId=94,UserId=407,IsDeleted=False,Status='Processing',
            ResultJson='BE result',RelatedEntityId=77,NotifyOnFinish=True,ProgressPercentage=0))
        async def execute(query):return conn.execute(query)
        with use_principal(who()):
            repo=AIRepository(SimpleNamespace(execute=execute),{'AI_Tasks':task})
            repository_patch(monkeypatch,erp_tasks,repo)
            append=AsyncMock();monkeypatch.setattr(erp_tasks,'append_version',append)
            @ai_outputs.tracked_workflow('Aggregate')
            async def workflow(request):
                if fails:raise ValueError('test failure')
                return {'request':request,'template':'BC','template_name':'Báo cáo','template_id':2,'template_version':3}
            token=ai_session.erp_task_id.set(554)
            try:
                if fails:
                    with pytest.raises(ValueError):await workflow('Báo cáo')
                else:
                    result=await workflow('Báo cáo')
                    assert result['ai_task_id']==554
                    append.assert_awaited_once()
            finally:ai_session.erp_task_id.reset(token)
        row=conn.execute(select(task)).mappings().one()
        assert row['Status']=='Processing' and row['ResultJson']=='BE result'
        assert row['RelatedEntityId']==77 and row['NotifyOnFinish'] is True
        assert row['ProgressPercentage']==(1 if fails else 100)

@pytest.mark.parametrize('kind,owner,member,deny,tenant,expected',[
    ('Common',None,False,False,94,True),('Personal',407,False,False,94,True),
    ('Personal',408,False,False,94,False),('Group',None,True,False,94,True),
    ('Group',None,False,False,94,False),('Common',None,False,True,94,False),
    ('Common',None,False,False,95,False)])
async def test_partition_access(kind,owner,member,deny,tenant,expected):
    metadata=MetaData()
    def table(name,columns):
        return Table(name,metadata,*[Column(n,t,primary_key=(n=='Id')) for n,t in columns.items()])
    audit=dict(Id=Integer,TenantId=Integer,IsDeleted=Boolean)
    doc=table('AI_KnowledgeDocuments',{**audit,'PartitionId':Integer,'FolderId':Integer,'FieldCode':String,
        'DocumentTypeCode':String,'SecurityLevelCode':String,'CreatorUserId':Integer})
    part=table('AI_DataPartitions',{**audit,'IsActive':Boolean,'Status':String,'PartitionType':String,
        'OwnerUserId':Integer,'WorkDepartmentId':Integer,'UserGroupId':Integer})
    group=table('AI_UserGroups',{**audit,'IsActive':Boolean,'Status':String})
    members=table('AI_UserGroupMembers',{**audit,'UserGroupId':Integer,'UserId':Integer,'LeftTime':DateTime})
    rules=table('AI_PartitionPermissions',{**audit,'SubjectType':String,'UserId':Integer,'RoleId':Integer,
        'UserGroupId':Integer,'WorkDepartmentId':Integer,'PartitionId':Integer,'FolderId':Integer,'DocumentId':Integer,
        'FieldCode':String,'DocumentTypeCode':String,'MaxSecurityLevelCode':String,'PermissionLevel':String,'IsDeny':Boolean})
    employee=table('Dms_Employee',{**audit,'UserId':Integer,'WorkDepartmentId':Integer})
    engine=create_engine('sqlite://');metadata.create_all(engine)
    with engine.begin() as conn:
        conn.execute(doc.insert().values(Id=422,TenantId=tenant,IsDeleted=False,PartitionId=1,CreatorUserId=408))
        conn.execute(part.insert().values(Id=1,TenantId=tenant,IsDeleted=False,IsActive=True,Status='Active',
            PartitionType=kind,OwnerUserId=owner,UserGroupId=3 if kind=='Group' else None))
        conn.execute(group.insert().values(Id=3,TenantId=tenant,IsDeleted=False,IsActive=True,Status='Active'))
        if member:conn.execute(members.insert().values(Id=1,TenantId=tenant,IsDeleted=False,UserGroupId=3,UserId=407))
        if deny:conn.execute(rules.insert().values(Id=1,TenantId=tenant,IsDeleted=False,SubjectType='User',
            UserId=407,PartitionId=1,PermissionLevel='ReadOnly',IsDeny=True))
        async def execute(query):return conn.execute(query)
        with use_principal(who()):
            repo=AIRepository(SimpleNamespace(execute=execute),{t.name:t for t in metadata.tables.values()})
            rows=await repo.rows('KnowledgeDocuments',owned=True)
            assert bool(rows)==expected


def test_legacy_be_minio_reference(monkeypatch):
    from app.core.config import get_settings
    from app.services.object_storage import parse_reference
    settings=get_settings()
    monkeypatch.setattr(settings,'be_minio_endpoint','https://s3.example.test')
    monkeypatch.setattr(settings,'be_minio_buckets',['be-bucket'])
    assert parse_reference('minio:knowledge/94/1/source.docx')==('be-bucket','knowledge/94/1/source.docx')
    assert parse_reference('minio://other/path.docx')==('other','path.docx')


def test_inline_article_text_never_disappears():
    from app.rag.chunking import Chunker,approximate_token_counter
    from app.core.config import get_settings
    text='### QUY CHẾ LÀM VIỆC\n\nĐiều 1. Thời gian làm việc: Buổi sáng từ 7 giờ 30 đến 11 giờ 30.\n\nĐiều 2. Trách nhiệm: Cán bộ thực hiện đúng nội quy.'
    cfg=get_settings().model_copy(update={'chunk_legal_structure':True})
    chunks=Chunker(cfg,token_counter=approximate_token_counter).split(text)
    assert '7 giờ 30' in '\n'.join(c.text for c in chunks)
    assert 'đúng nội quy' in '\n'.join(c.text for c in chunks)


async def test_report_exports_allow_missing_number(monkeypatch,tmp_path):
    from app.agents.nodes import drafting,report
    from app.core.config import get_settings
    monkeypatch.setattr(get_settings(),'output_dir',str(tmp_path))
    captured=[]
    def build(payload,path,template):captured.append(payload.meta);return path
    monkeypatch.setattr(drafting,'build_docx',build)
    monkeypatch.setattr(report,'build_docx',build)
    monkeypatch.setattr(report.TemplateRepository,'get',AsyncMock(return_value=None))
    base={'request':'Báo cáo test','inputs':{},'params':{'ky':'2026-10','thang':10,'nam':2026},
          'sections':[],'data':{'queries':[]}}
    draft=await drafting.export_node({**base,'template':{'ma_template':'BC','ten_bao_cao':'Báo cáo','fields':{}}})
    aggregate=await report.export_node(base)
    assert draft['so_ky_hieu']==aggregate['so_ky_hieu']==''
    assert all(meta['so_ky_hieu']=='' for meta in captured)
