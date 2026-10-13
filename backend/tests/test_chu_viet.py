"""Chữ tiếng Việt từ SQL tới file .pptx: không mất dấu, không tên trường CSDL, file mở được trên Windows."""

from __future__ import annotations

import unicodedata
from io import BytesIO

import sqlglot
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Inches

from app.agents.nodes.presentation import _MAT_DAU
from app.documents import slide_deck as sd
from app.documents.nhan_cot import lam_sach_ket_qua, nhan_hop_le
from app.documents.pptx_svg import them_png_du_phong
from app.documents.query_result import table
from app.services.erp_sql_policy import _chuoi_unicode


# ------------------------------------------------------------------ SQL -- #
def test_chuoi_co_dau_trong_sql_thanh_n_prefix():
    sql = ("SELECT CASE WHEN a = 1 THEN 'Tồn cuối kỳ' ELSE N'Rời khỏi' END AS x, 'abc' AS y "
           "FROM t WHERE b LIKE '%đ%' AND c = 'it''s'")
    out = sqlglot.parse_one(sql, read="tsql").transform(_chuoi_unicode).sql(dialect="tsql")
    assert "N'Tồn cuối kỳ'" in out and "N'%đ%'" in out and "N'Rời khỏi'" in out
    assert "'abc'" in out and "N'abc'" not in out and "'it''s'" in out   # chuỗi ASCII giữ nguyên


# ------------------------------------------------------- nhãn cột, NFC -- #
def _ket_qua(columns, rows):
    return {"query_id": "q", "question": "Biến động nhân sự", "columns": columns, "rows": rows,
            "row_count": len(rows), "truncated": False, "schema_version": "s", "policy_version": "p",
            "source_tables": [], "warnings": []}


def test_khong_ten_truong_nao_len_tieu_de_cot():
    nfd = unicodedata.normalize("NFD", "Vào làm")
    r = _ket_qua([{"name": "soluong", "data_type": "int", "label": "soluong", "unit": "nguoi"},
                  {"name": "loai_chuyen_dong_xyz", "data_type": "nvarchar", "label": None, "unit": None},
                  {"name": "thangnam", "data_type": "nvarchar", "label": "", "unit": None}],
                 [[3, nfd, "2026-06"]])
    lam_sach_ket_qua(r)
    labels = [c["label"] for c in r["columns"]]
    # Chốt chặn không đoán/ánh xạ: nhãn tên trường hoặc rỗng -> nhãn trung tính (Qwen đặt ở bước trước).
    assert labels == ["Cột 1", "Cột 2", "Cột 3"]
    assert r["columns"][0]["unit"] == ""                      # "nguoi" không dấu: bỏ, không in sai
    assert r["rows"][0][1] == "Vào làm" and r["rows"][0][1] != nfd   # về NFC


def test_bang_word_va_slide_khong_lui_ve_ten_truong():
    r = _ket_qua([{"name": "soluongnhansu", "data_type": "int", "label": None, "unit": None},
                  {"name": "tenphongban", "data_type": "nvarchar", "label": None, "unit": None}],
                 [[5, "Phòng A"]])
    _, data = table(r)
    assert "soluongnhansu" not in " ".join(data["columns"]) and "tenphongban" not in " ".join(data["columns"])
    parts = sd.query_parts({"queries": [_ket_qua(
        [{"name": "soluong", "data_type": "int", "label": "soluong", "unit": None},
         {"name": "phongban", "data_type": "nvarchar", "label": "phongban", "unit": None}], [[5, "Phòng A"]])]})
    ds, = parts.datasets
    chu = " ".join([ds.title, ds.short, ds.group_label, *ds.table_columns, *(n for n, _ in ds.series), *ds.facts])
    assert "soluong" not in chu and "phongban" not in chu


def test_nhan_bieu_do_nfc():
    from app.documents.charts import dataset_chart

    nfd = unicodedata.normalize("NFD", "Phòng Kế toán")
    assert dataset_chart([nfd], [("Số", [1])])[1:4] == b"PNG"


# --------------------------------------------------------- SVG -> PNG -- #
_SVG = (b'<svg xmlns="http://www.w3.org/2000/svg" width="100" height="50" viewBox="0 0 100 50">'
        b'<rect width="100" height="50" fill="#2563EB"/></svg>')


def _pptx_svg_thuan() -> bytes:
    """Đúng cấu trúc Presenton ghi: <a:blip> không r:embed, chỉ có svgBlip trong extLst."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    png = BytesIO()
    from PIL import Image
    Image.new("RGB", (10, 10), "white").save(png, "PNG")
    pic = slide.shapes.add_picture(BytesIO(png.getvalue()), Inches(1), Inches(1), Inches(2), Inches(1))
    svg_part = Part(PackURI("/ppt/media/hinh-1.svg"), "image/svg+xml", slide.part.package, _SVG)
    rid = slide.part.relate_to(svg_part, RT.IMAGE)
    blip = pic._element.find(".//" + qn("a:blip"))
    del blip.attrib[qn("r:embed")]
    ext_lst = blip.makeelement(qn("a:extLst"), {})
    ext = ext_lst.makeelement(qn("a:ext"), {"uri": "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"})
    svg = ext.makeelement("{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip", {qn("r:embed"): rid})
    ext.append(svg)
    ext_lst.append(ext)
    blip.append(ext_lst)
    out = BytesIO()
    prs.save(out)
    return out.getvalue()


def test_anh_svg_thuan_duoc_them_png_du_phong():
    goc = _pptx_svg_thuan()
    blip = Presentation(BytesIO(goc)).slides[0]._element.find(".//" + qn("a:blip"))
    assert blip.get(qn("r:embed")) is None                     # tái hiện đúng lỗi của Presenton
    sua = them_png_du_phong(goc)
    slide = Presentation(BytesIO(sua)).slides[0]
    blip = slide._element.find(".//" + qn("a:blip"))
    rid = blip.get(qn("r:embed"))
    assert rid and slide.part.related_part(rid).content_type == "image/png"
    # SVG vẫn giữ trong extLst cho bản PowerPoint đọc được nó.
    assert blip.find(".//{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip") is not None


def test_file_khong_co_svg_thuan_giu_nguyen():
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    out = BytesIO()
    prs.save(out)
    assert them_png_du_phong(out.getvalue()) == out.getvalue()


# ------------------------------------------------- phát hiện mất dấu -- #
def test_nhan_dien_chu_mat_dau():
    assert _MAT_DAU.search("2026-09 · T?n cu?i k?")
    assert _MAT_DAU.search("B�o c�o")
    assert not _MAT_DAU.search("Tại sao số liệu giảm?")
    assert not _MAT_DAU.search("Tồn cuối kỳ")
