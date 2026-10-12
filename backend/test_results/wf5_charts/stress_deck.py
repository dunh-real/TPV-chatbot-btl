"""Dựng slide thật cho từng ca của label_stress (dữ liệu đã gắn nhãn) qua cả hai đường và kiểm như eval_charts.

    python stress_deck.py <runs/stressN> [S01,S05,...]
Đọc <run>/cases/<id>.json, ghi <run>/decks/<id>/{presenton.pptx, local.pptx, presenton_slide_NN.png} và
<run>/decks/result.json. Không gọi lại bước đặt nhãn (nhãn đã có): đo riêng phần dựng slide.
"""

from __future__ import annotations

import asyncio
import json
import re
import sys
import time
import unicodedata
from pathlib import Path

PARAMS = {"ky": "2026-09", "thang": 9, "nam": 2026}


def texts(path: Path) -> list[str]:
    from pptx import Presentation

    out = []
    for slide in Presentation(str(path)).slides:
        stack = list(slide.shapes)
        while stack:
            shp = stack.pop()
            if shp.shape_type == 6:
                stack.extend(shp.shapes)
            if shp.has_text_frame and shp.text_frame.text.strip():
                out.append(shp.text_frame.text)
            if getattr(shp, "has_table", False) and shp.has_table:
                out += [c.text for row in shp.table.rows for c in row.cells if c.text.strip()]
    return out


async def main() -> None:
    sys.path.insert(0, str(Path(__file__).parent))
    from eval_charts import pdf_to_png, pictures_by_slide, svg_only_blips

    from app.agents.nodes import presentation as pr
    from app.agents.nodes.presentation import _MAT_DAU
    from app.documents.nhan_cot import la_ten_truong
    from app.services.presenton import close_presenton, get_presenton

    root = Path(sys.argv[1])
    ids = sys.argv[2].split(",") if len(sys.argv) > 2 else sorted(p.stem for p in (root / "cases").glob("*.json"))
    results = []
    for cid in ids:
        q = json.loads((root / "cases" / f"{cid}.json").read_text(encoding="utf-8"))
        out = root / "decks" / cid
        out.mkdir(parents=True, exist_ok=True)
        raw_names = {c["name"] for c in q["columns"] if len(c["name"]) >= 4}
        state = {"request": q["question"], "data": {"queries": [q]}, "params": PARAMS, "assumptions": [],
                 "inputs": {"don_vi_trinh_bay": "Đơn vị thử"}}
        rec: dict = {"id": cid, "bang": q["question"]}
        try:
            t0 = time.monotonic()
            res = await pr._render_direct(dict(state), out / "presenton.pptx")
            rec["engine"], rec["seconds"] = res.get("engine"), round(time.monotonic() - t0, 1)
            if res.get("engine") == "presenton":
                dc = res.get("deck_check") or {}
                rec.update(deck_ok=dc.get("ok"), charts=f"{dc.get('charts_ok')}/{dc.get('charts_total')}",
                           summary=res.get("summary_source"), svg_only=svg_only_blips(Path(res["output_path"])))
                try:
                    pdf_to_png(await get_presenton().export_deck(res["presentation_id"], "pdf"), out, "presenton")
                except Exception as exc:  # noqa: BLE001
                    rec["pdf_error"] = str(exc)[:120]
            else:
                rec["fallback_reason"] = (res.get("assumptions") or [""])[-1][:200]
            val = await pr.verify_node({**state, **res})
            rec["validation"] = val["validation"].get("status")
            rec["issues"] = [i.get("type") for i in val["validation"].get("issues", [])]
            loc = await pr._render_local(dict(state), out / "local.pptx", RuntimeError("thử đường lùi"))
            vloc = await pr.verify_node({**state, **loc})
            rec["local_validation"] = vloc["validation"].get("status")
            rec["local_charts"] = sum(len(p) for p in pictures_by_slide(Path(loc["output_path"])))
            bad = []
            for f in (out / "presenton.pptx", out / "local.pptx"):
                for t in texts(f) if f.exists() else []:
                    if la_ten_truong(t) or _MAT_DAU.search(t) or t != unicodedata.normalize("NFC", t) or any(
                            re.search(rf"\b{re.escape(n)}\b", t) for n in raw_names if not re.search(r"[^\x00-\x7f]", n)):
                        bad.append(f"{f.stem}: {t[:80]}")
            rec["bad_text"] = bad[:5]
        except Exception as exc:  # noqa: BLE001
            rec["error"] = f"{type(exc).__name__}: {exc}"[:300]
        results.append(rec)
        print(json.dumps(rec, ensure_ascii=False), flush=True)
    await close_presenton()
    (root / "decks" / "result.json").write_text(json.dumps(results, ensure_ascii=False, indent=1), encoding="utf-8")


if __name__ == "__main__":
    asyncio.run(main())
