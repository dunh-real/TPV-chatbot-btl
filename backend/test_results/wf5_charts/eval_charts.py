"""Kiểm thử workflow 5 sau khi thống nhất biểu đồ (app.documents.charts) cho cả hai đường dựng slide.

10 câu hỏi của report_pipeline, số liệu SQL thật từ ERP (login chỉ đọc), đưa vào backend đúng dạng kết
quả Text-to-SQL (`data["queries"]`). Mỗi câu dựng HAI lần:

    presenton  đường mặc định: biểu đồ PNG do code vẽ, tải lên Presenton, gắn vào ô biểu đồ của layout
    local      đường lùi pptx_builder: cùng hàm vẽ, slide tự dựng

Đánh giá mỗi câu: engine thực tế, đối chiếu deck (deck_check), đối chiếu số trên file (verify_node), số
biểu đồ vẽ ra so với số bộ dữ liệu, SỐ TRÊN BIỂU ĐỒ đối chiếu ĐỘC LẬP với các dòng SQL gốc (không qua
code dựng biểu đồ), có ảnh biểu đồ trong file .pptx của từng đường, thời gian. Ảnh biểu đồ và PNG từng
slide được lưu để xem bằng mắt.

Không cần vLLM: PRESENTON_SUMMARY=false, phần chữ nhận xét do code viết.
Chạy qua run_eval.sh (dựng chuỗi kết nối từ secrets, không in ra).
"""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import io
import json
import os
import sys
import time
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

PARAMS = {"ky": "2026-09", "thang": 9, "nam": 2026}
TEN_DON_VI = {94: "BQP1", 63: "Công ty TPV", None: "Toàn hệ thống"}


def _kieu(py_type: Any) -> str:
    if py_type is bool:
        return "bit"
    if py_type is int:
        return "int"
    if py_type is float:
        return "float"
    if py_type is Decimal:
        return "decimal"
    if py_type in (date, datetime):
        return "datetime"
    return "nvarchar"


def _gia_tri(value: Any) -> Any:
    # Giống as_dict của Text-to-SQL: Decimal thành chuỗi để không mất chính xác.
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def fetch(conn, key: str, spec: dict[str, Any]) -> dict[str, Any]:
    from app.documents.nhan_cot import don_vi_mac_dinh, nhan_tu_ten

    cur = conn.cursor()
    cur.execute(spec["sql"])
    desc = cur.description
    rows = [[_gia_tri(v) for v in r] for r in cur.fetchall()]
    columns = []
    for d in desc:
        label = nhan_tu_ten(d[0]) or d[0]
        columns.append({"name": d[0], "data_type": _kieu(d[1]), "label": label, "unit": don_vi_mac_dinh(label)})
    # Đủ trường của app.schemas.erp.QueryResult: đường lùi dựng bảng từ chính schema đó.
    return {"query_id": key, "question": spec["title"], "columns": columns, "rows": rows,
            "row_count": len(rows), "truncated": False, "schema_version": "eval", "policy_version": "eval",
            "source_tables": [], "warnings": [], "dataset_key": key}


def _so(v: Any) -> float | None:
    try:
        return float(Decimal(str(v)))
    except Exception:  # noqa: BLE001
        return None


def check_chart_values(ds, query: dict[str, Any]) -> list[str]:
    """Số trên biểu đồ đối chiếu với các dòng SQL gốc - tra theo NHÃN, không dùng lại logic dựng biểu đồ."""
    from app.documents.slide_deck import _nhan

    cols, rows = query["columns"], query["rows"]
    problems: list[str] = []
    for series_name, values in ds.series:
        idx = [i for i, c in enumerate(cols) if _nhan(c["label"]) == series_name]
        if not idx:
            problems.append(f"chuỗi {series_name!r} không ứng với cột nào")
            continue
        j = idx[0]
        grouped = ds.grouped_others and ds.categories[-1].endswith("khác")
        cats = ds.categories[:-1] if grouped else ds.categories
        for cat, v in zip(cats, values):
            khop = [r for r in rows if any(_nhan(x) == cat for x in r if isinstance(x, str))
                    or (cat == "Chưa xác định" and None in r)]
            gia_tri = [_so(r[j]) or 0.0 for r in khop]
            if not khop:
                problems.append(f"{cat!r}: không có dòng SQL nào")
            elif float(v) not in gia_tri:
                problems.append(f"{cat!r}/{series_name}: biểu đồ {v} != SQL {gia_tri}")
        if grouped:
            tong_sql = sum(_so(r[j]) or 0.0 for r in rows)
            if abs(sum(values) - tong_sql) > 1e-6:
                problems.append(f"{series_name}: tổng biểu đồ {sum(values)} != tổng SQL {tong_sql}")
    return problems


def pictures_by_slide(path: Path) -> list[list[bytes]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    out = []
    for slide in Presentation(str(path)).slides:
        pics = []
        stack = list(slide.shapes)
        while stack:
            shp = stack.pop()
            if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
                stack.extend(shp.shapes)
            elif shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    pics.append(shp.image.blob)
                except ValueError:          # ảnh liên kết/không nhúng: vẫn là một ảnh trên slide
                    pics.append(b"")
        out.append(pics)
    return out


def svg_only_blips(path: Path) -> int:
    """Ảnh SVG không có PNG dự phòng - PowerPoint trên Windows không hiển thị/đòi sửa file."""
    from pptx import Presentation
    from pptx.oxml.ns import qn

    prs = Presentation(str(path))
    parts = [s.part for s in prs.slides] + [m.part for m in prs.slide_masters]
    return sum(1 for p in parts for b in p._element.iter(qn("a:blip")) if not b.get(qn("r:embed")))


def pdf_to_png(pdf: bytes, out_dir: Path, prefix: str) -> int:
    import pymupdf

    doc = pymupdf.open(stream=pdf, filetype="pdf")
    for i, page in enumerate(doc, start=1):
        page.get_pixmap(dpi=60).save(str(out_dir / f"{prefix}_slide_{i:02d}.png"))
    return doc.page_count


async def run_one(q: dict[str, Any], datasets: dict[str, Any], conn, out_root: Path) -> dict[str, Any]:
    from app.agents.nodes import presentation as pr
    from app.documents import slide_deck as sd
    from app.documents.charts import chart_for
    from app.services.presenton import get_presenton

    out = out_root / q["id"]
    out.mkdir(parents=True, exist_ok=True)
    queries = [fetch(conn, k, datasets[k]) for k in q["datasets"]]
    data = {"queries": [{k: v for k, v in x.items() if k != "dataset_key"} for x in queries]}
    (out / "data.json").write_text(json.dumps(queries, ensure_ascii=False, indent=1, default=str), encoding="utf-8")

    dsets = sd.build_datasets(data, PARAMS)
    index_of = {f"bang_{i + 1}": i for i in range(len(queries))}
    chart_problems = {d.key: check_chart_values(d, data["queries"][index_of[d.key]]) for d in dsets}
    for d in dsets:
        (out / f"chart_{d.key}.png").write_bytes(chart_for(d))

    state = {"request": q["question"], "data": data, "params": PARAMS, "assumptions": [],
             "inputs": {"don_vi_trinh_bay": TEN_DON_VI.get(q["tenant"], "Đơn vị")}}
    rec: dict[str, Any] = {"id": q["id"], "question": q["question"], "datasets": len(dsets),
                           "rows": [len(x["rows"]) for x in queries], "chart_value_problems": chart_problems}

    # --- đường Presenton
    t0 = time.monotonic()
    res = await pr._render_direct(dict(state), out / "presenton.pptx")
    rec_p: dict[str, Any] = {"engine": res.get("engine"), "seconds": round(time.monotonic() - t0, 1),
                             "slides": res.get("slide_count"), "assumptions": res.get("assumptions", [])}
    if res.get("engine") == "presenton":
        dc = res.get("deck_check") or {}
        rec_p.update(deck_ok=dc.get("ok"), charts_ok=dc.get("charts_ok"), charts_total=dc.get("charts_total"),
                     deck_issues=dc.get("issues", [])[:5], timings=res.get("timings"),
                     presentation_id=res.get("presentation_id"))
        val = await pr.verify_node({**state, **res})
        rec_p["validation"] = val["validation"].get("status")
        rec_p["validation_issues"] = val["validation"].get("issues", [])[:5]
        pics = pictures_by_slide(Path(res["output_path"]))
        rec_p["svg_only_blips"] = svg_only_blips(Path(res["output_path"]))
        chart_idx = [i for i, p in enumerate(res["brief"].split("\n\n---\n\n")) if p.startswith("[chart]")]
        rec_p["chart_slides_with_picture"] = sum(1 for i in chart_idx if i < len(pics) and pics[i])
        rec_p["chart_slides"] = len(chart_idx)
        try:
            pdf = await get_presenton().export_deck(res["presentation_id"], "pdf")
            rec_p["pdf_pages"] = pdf_to_png(pdf, out, "presenton")
        except Exception as exc:  # noqa: BLE001
            rec_p["pdf_error"] = str(exc)[:200]
    rec["presenton"] = rec_p

    # --- đường lùi
    t0 = time.monotonic()
    loc = await pr._render_local(dict(state), out / "local.pptx", RuntimeError("kiểm thử đường lùi"))
    rec_l: dict[str, Any] = {"engine": loc.get("engine"), "seconds": round(time.monotonic() - t0, 1),
                             "slides": loc.get("slide_count")}
    val = await pr.verify_node({**state, **loc})
    rec_l["validation"] = val["validation"].get("status")
    rec_l["validation_issues"] = val["validation"].get("issues", [])[:5]
    blobs = [b for pics in pictures_by_slide(Path(loc["output_path"])) for b in pics]
    expected = {hashlib.sha256(chart_for(d)).hexdigest() for d in dsets}
    rec_l["charts_in_file"] = len(blobs)
    rec_l["charts_same_as_chart_for"] = sum(hashlib.sha256(b).hexdigest() in expected for b in blobs)
    rec["local"] = rec_l
    (out / "result.json").write_text(json.dumps(rec, ensure_ascii=False, indent=1), encoding="utf-8")
    return rec


async def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--questions", required=True, help="thư mục report_pipeline (chứa questions.py)")
    ap.add_argument("--out", required=True)
    ap.add_argument("--only", default="")
    args = ap.parse_args()

    sys.path.insert(0, args.questions)
    import pyodbc
    from questions import DATASETS, QUESTIONS

    from app.services.presenton import close_presenton

    out_root = Path(args.out)
    out_root.mkdir(parents=True, exist_ok=True)
    only = {x for x in args.only.split(",") if x}
    conn = pyodbc.connect(os.environ["T2S_ODBC"], readonly=True, timeout=30)
    results = []
    try:
        for q in QUESTIONS:
            if only and q["id"] not in only:
                continue
            started = time.monotonic()
            try:
                rec = await run_one(q, DATASETS, conn, out_root)
            except Exception as exc:  # noqa: BLE001 - một câu hỏng không dừng cả lượt
                rec = {"id": q["id"], "error": f"{type(exc).__name__}: {exc}"[:500]}
            rec["total_seconds"] = round(time.monotonic() - started, 1)
            results.append(rec)
            print(json.dumps({k: rec.get(k) for k in ("id", "error", "total_seconds")}, ensure_ascii=False),
                  flush=True)
    finally:
        conn.close()
        await close_presenton()
    (out_root / "results.json").write_text(json.dumps(results, ensure_ascii=False, indent=1), encoding="utf-8")


if __name__ == "__main__":
    asyncio.run(main())
